"""Extrait d'un deck de référence (« Trade in review » FR) les éléments statiques réutilisés par le format « revue mensuelle France » :
logos des marques, drapeau, bloc Stellantis de la couverture et diapositives de l'annexe « Trade-in CTA » (captures des sites + légendes + positions).

Usage : python3 scripts/extract-revue-assets.py "<deck.pptx>" <dossier de sortie, ex. data/presentations/assets/revue_fr>
Les images ne sont PAS versionnées dans ce repo (captures de sites de marque, logos) : elles vivent dans le repo de données privé, sous
presentations/assets/revue_fr/ ; le générateur les lit là (Donnees.hist.parent / "presentations" / "assets" / "revue_fr")."""

import json
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

EMU = 914400
# diapositive « Traffic » de chaque marque dans le deck de référence (le logo de la marque y est en haut, au centre)
MARQUES = {"PEUGEOT": 16, "CITROEN": 22, "DS": 28, "OPEL": 34, "FIAT": 40, "ALFA ROMEO": 46, "JEEP": 52, "LANCIA": 58}
ANNEXE = range(64, 73)        # « Trade-in CTA on main / brand website »


def pouces(v):
    return round(v / EMU, 3)


def sauve(sh, dossier, nom):
    blob = sh.image.blob
    chemin = dossier / f"{nom}.png"
    chemin.write_bytes(blob)
    try:                      # PNG optimisé : les captures de site pèsent plusieurs Mo
        im = Image.open(chemin).convert("RGBA")
        g, h, d, b_ = sh.crop_left, sh.crop_top, sh.crop_right, sh.crop_bottom      # recadrage appliqué dans le deck (l'image d'origine est souvent plus grande)
        if any((g, h, d, b_)):
            im = im.crop((int(im.width * g), int(im.height * h), int(im.width * (1 - d)), int(im.height * (1 - b_))))
        if im.width > 1800:
            im = im.resize((1800, int(im.height * 1800 / im.width)))
        im.save(chemin, optimize=True)
    except Exception:
        pass
    return chemin.name


def main(deck, sortie):
    prs = Presentation(deck)
    slides = list(prs.slides)
    dossier = Path(sortie)
    dossier.mkdir(parents=True, exist_ok=True)
    manifeste = {"logos": {}, "drapeau": None, "bloc_stellantis": None, "annexe": []}
    # logos : l'image carrée en haut, entre x = 5 et 9.5 po
    for marque, n in MARQUES.items():
        for sh in slides[n - 1].shapes:
            if sh.shape_type == 13 and sh.top < 0.4 * EMU and 5.0 * EMU < sh.left < 9.5 * EMU and sh.width < 2.6 * EMU:
                manifeste["logos"][marque] = {"fichier": sauve(sh, dossier, f"logo_{marque.lower().replace(' ', '_')}"), "x": pouces(sh.left), "y": pouces(sh.top), "w": pouces(sh.width), "h": pouces(sh.height)}
                break
    # drapeau (en haut à droite) et bloc Stellantis de la couverture
    for sh in slides[15].shapes:
        if sh.shape_type == 13 and sh.left > 10.2 * EMU and sh.top < 0.5 * EMU:
            manifeste["drapeau"] = {"fichier": sauve(sh, dossier, "drapeau_fr"), "x": pouces(sh.left), "y": pouces(sh.top), "w": pouces(sh.width), "h": pouces(sh.height)}
    for sh in slides[0].shapes:
        if sh.shape_type == 13:
            nom = "couverture_drapeau" if sh.width < 1.2 * EMU else "couverture_stellantis"
            manifeste["couverture_" + ("drapeau" if sh.width < 1.2 * EMU else "stellantis")] = {"fichier": sauve(sh, dossier, nom), "x": pouces(sh.left), "y": pouces(sh.top), "w": pouces(sh.width), "h": pouces(sh.height)}
    # annexe : toutes les images de contenu + les légendes (zones de texte), avec leurs positions
    for n in ANNEXE:
        s = slides[n - 1]
        titre = s.shapes.title.text if s.shapes.title is not None else ""
        elements = []
        for i, sh in enumerate(s.shapes):
            if sh.shape_type == 13:
                if sh.top < 1.0 * EMU and sh.left > 10.0 * EMU and sh.width < 1.6 * EMU:
                    continue          # drapeau : posé par le générateur (le logo de la marque, lui, est repris tel quel)
                elements.append({"type": "image", "fichier": sauve(sh, dossier, f"annexe_{n}_{i}"), "x": pouces(sh.left), "y": pouces(sh.top), "w": pouces(sh.width), "h": pouces(sh.height)})
            elif sh.has_text_frame and sh.text_frame.text.strip() and not sh.is_placeholder:
                r = [r for p in sh.text_frame.paragraphs for r in p.runs]
                elements.append({"type": "texte", "texte": sh.text_frame.text.strip(), "x": pouces(sh.left), "y": pouces(sh.top), "w": pouces(sh.width), "h": pouces(sh.height),
                                 "taille": (r[0].font.size.pt if r and r[0].font.size else 12), "gras": bool(r and r[0].font.bold)})
        manifeste["annexe"].append({"diapositive": n, "titre": titre, "elements": elements})
    (dossier / "manifeste.json").write_text(json.dumps(manifeste, ensure_ascii=False, indent=1), encoding="utf-8")
    print("logos:", sorted(manifeste["logos"]), "| annexe:", len(manifeste["annexe"]), "diapositives |", sum(f.stat().st_size for f in dossier.iterdir()) // 1024, "Ko")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
