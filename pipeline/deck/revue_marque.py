"""Revue mensuelle France : diapositives par marque « Traffic » (16, 22, …) et « YTD Leads Trends » (17, 23, …)."""

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from . import graphiques as g
from .revue_commun import evol, graphique_annuel, mois_label, nb, pct, slide_revue

C_TUILE = "BBDEFB"
VERT, ROUGE, BLEU_EVOL = "2E9E4F", "D93025", "1A73E8"


def _fmt(v, n=1):
    return "n/a" if v is None else f"{v:,.{n}f}".replace(",", " ")


def _d(cur, ref):
    """Variation en % avec signe (None si base nulle)."""
    return evol(cur, ref) if (cur is not None and ref) else None


def _txt_evol(e):
    return "n/a" if e is None else ("▲ " if e > 0 else "▼ " if e < 0 else "") + f"{e:+.1f}%".replace("+", "")


def _couleur(e):
    return "6B7280" if e is None else (VERT if e > 0 else ROUGE if e < 0 else "6B7280")


def _carre(slide, x, y, w, h, etiquette, valeur, evols, taille_valeur=18):
    """Tuile bleu clair du trafic : libellé, valeur, variations (liste de (libellé, %))."""
    f = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    f.fill.solid(); f.fill.fore_color.rgb = g.rgb(C_TUILE); f.line.fill.background(); f.shadow.inherit = False
    g.texte(slide, x + 0.08, y + 0.04, w - 0.16, 0.22, etiquette, taille=8, couleur="595959")
    g.texte(slide, x + 0.08, y + 0.22, w - 0.16, 0.4, valeur, taille=taille_valeur, couleur="000000")
    for i, (lib, e) in enumerate(evols):
        g.texte(slide, x + 0.08 + i * (w - 0.16) / max(1, len(evols)), y + h - 0.3, (w - 0.16) / max(1, len(evols)), 0.26,
                f"{_txt_evol(e)} {lib}", taille=8, couleur=_couleur(e))


def _ligne(slide, x1, y1, x2, y2, fleche=False):
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.oxml.ns import qn
    from lxml import etree
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = g.rgb("000000"); c.line.width = Pt(0.75)
    if fleche:
        ln = c.line._get_or_add_ln()
        etree.SubElement(ln, qn("a:tailEnd")).set("type", "triangle")
    return c


def _taux(a, b):
    return None if (a is None or not b) else a / b * 100


def _mois_cle(per):
    return f"{per.annee}-{per.indice:02d}"


def _cles(per):
    """Clés « AAAA-MM » des mois de la période (un seul mois, ou les 3 d'un trimestre)."""
    return [f"{a}-{m:02d}" for a, m in per.mois()]


def _entree(ctx, marque, per):
    """{canal: sessions} des arrivées « Trade In Entry » du mois, sommées sur les sites de la marque ; {} si absent."""
    out = {}
    for nom, m, _ in ctx.d.sites:
        if m != marque:
            continue
        for cle in _cles(per):
            for canal, n in ((((ctx.d._site_json(nom) or {}).get("entreeMonth") or {}).get(cle) or {}).get("canaux") or {}).items():
                out[canal] = out.get(canal, 0) + n
    return out


def _entree_origines(ctx, marque, per):
    """(entrées depuis le site de la marque, entrées directes « Trade In Entry ») du mois ou du trimestre, sommées sur les sites de la marque ; None si absent (champs « marque » et « total » de entreeMonth)."""
    marque_n = direct_n = 0
    vu = False
    for nom, m, _ in ctx.d.sites:
        if m != marque:
            continue
        for cle in _cles(per):
            e = ((ctx.d._site_json(nom) or {}).get("entreeMonth") or {}).get(cle)
            if e and e.get("marque") is not None:
                marque_n += e.get("marque") or 0
                direct_n += e.get("total") or 0
                vu = True
    return (marque_n, direct_n) if vu and (marque_n + direct_n) else None


def _site_principal(ctx, marque, per):
    """(sessions du site principal, sessions du site de reprise) du mois, sommées sur les sites de la marque (data/<site>.json : propSessions.month = propriété GA4 entière, comme Looker / repriseMonth) ; None si absent."""
    a = r = None
    for nom, m, _ in ctx.d.sites:
        if m != marque:
            continue
        j = ctx.d._site_json(nom) or {}
        vs = [((j.get("propSessions") or {}).get("month") or {}).get(cle) for cle in _cles(per)]       # « Visites site marque » de Looker = sessions de toute la propriété GA4 (build.py / looker_align.py)
        ws = [((j.get("repriseMonth") or {}).get(cle) or {}).get("sessions") for cle in _cles(per)]
        if all(v is not None for v in vs):                  # jamais de total partiel sur un trimestre
            a = (a or 0) + sum(vs)
        if all(w is not None for w in ws):
            r = (r or 0) + sum(ws)
    return a, r


def trafic(ctx, rv, marque):
    d = ctx.d
    nom = ctx.nom(marque)
    # tuiles du haut : « période précédente » = les mêmes nombres de jours juste avant (définition de Looker, comme la diapo Traffic & leads Overview) sur le trimestriel
    p_ses = rv.per.precedente_jours() if rv.trim else rv.prec
    cur, prec, n1 = (d.utilisateurs(p, marque) for p in (rv.per, p_ses, rv.n1))
    can, can_p, can_n = (d.canaux(p, marque) or {} for p in (rv.per, rv.prec, rv.n1))
    ses = cur["sessions"] if cur else (sum(can.values()) or None)
    if not ses:
        return None

    def val(u, k, repli=None):
        return u[k] if u else repli
    s_p = val(prec, "sessions", sum(can_p.values()) or None)
    s_n = val(n1, "sessions", sum(can_n.values()) or None)
    # tableau et anneau « Traffic landing on Trade In » = sessions « Trade In Entry » (arrivées sur la home page de la reprise) par canal, comme Looker :
    # data/<site>.json -> entreeMonth[mois].canaux (pipeline/ga4.py entree_directe_par_canal) ; repli sur les canaux de toutes les sessions si absent
    ent = [_entree(ctx, marque, p) for p in (rv.per, rv.prec, rv.n1)]
    if ent[0]:
        can, can_p, can_n = ent[0], ent[1] or {}, ent[2] or {}
    ij, hl = val(cur, "in_journey"), val(cur, "hot_leads")
    e_s = (_d(ses, s_p), _d(ses, s_n))
    e_ij = (_d(ij, val(prec, "in_journey")), _d(ij, val(n1, "in_journey")))
    e_hl = (_d(hl, val(prec, "hot_leads")), _d(hl, val(n1, "hot_leads")))
    pays_dd = "+".join(ctx.d.pays)
    for p_, u_ in ((rv.per, cur), (p_ses, prec), (rv.n1, n1)):
        if u_:
            for k_, nom_ in (("sessions", "sessions"), ("in_journey", "in_journey"), ("hot_leads", "ga_leads")):
                ctx.faits.ajoute(nom_, u_.get(k_), marque, p_.etiquette("en"), "traffic_deep_dive", perimetre=pays_dd)
    site, rep = _site_principal(ctx, marque, rv.per)
    site_p = _site_principal(ctx, marque, rv.prec)[0]
    site_n = _site_principal(ctx, marque, rv.n1)[0]
    if rv.trim:                      # un trimestre = UNE requête GA4 (Looker) : additionner les trois mois surévalue (GA4 dédoublonne), PEUGEOT PL T3 : 741 323 au lieu de 730 864
        vivant = [ctx.d.sessions_site_marque(p, marque) for p in (rv.per, p_ses, rv.n1)]
        site, site_p, site_n = (v if v is not None else r for v, r in zip(vivant, (site, site_p, site_n)))
    e_site = (_d(site, site_p), _d(site, site_n))
    t1, t2, cvr = _taux(ij, ses), _taux(hl, ij), _taux(hl, ses)
    ratio = _taux(rep if rep is not None else ses, site)

    # commentaires (règles)
    ordre = sorted(can, key=lambda c: -can[c])
    puces = []
    if e_s[1] is not None:
        tr = f"Traffic to trade-in websites {'increased' if e_s[1] > 0 else 'decreased'} YoY ({pct(e_s[1], 1, True)})"
        if e_site[1] is not None:
            tr += f" and {'outperformed' if e_s[1] > e_site[1] else 'underperformed'} {nom} main website traffic ({pct(e_site[1], 1, True)} YoY)"
        puces.append(tr + ".")
    else:
        puces.append(f"Trade-in website sessions: {nb(ses)}" + (f" ({pct(e_s[0], 1, True)} {rv.vs_prec})." if e_s[0] is not None else "."))
    if ratio is not None:
        puces.append(f"Ratio trade-in sessions / main website sessions: {ratio:.1f}%")
    if hl is not None and cvr is not None:
        p = f"Funnel: {nb(ses)} sessions, {nb(ij)} in journey, {nb(hl)} GA leads"
        if e_hl[1] is not None:
            p += f" ({pct(e_hl[1], 1, True)} YoY GA leads)"
        puces.append(p + f"; overall conversion rate {cvr:.1f}%.")
    else:
        puces.append("In Journey and GA leads: n/a (GA4 unavailable).")
    ct = sum(can.values())
    if ordre and ct:
        c0 = ordre[0]
        p = f"Main channel: {c0}, {can[c0] / ct * 100:.1f}% of trade-in traffic"
        e0 = _d(can[c0], can_n.get(c0)) if can_n else None
        if e0 is not None:
            p += f" ({pct(e0, 1, True)} YoY)"
        puces.append(p + ".")

    titre, sous = ("Traffic DEEP DIVE", f"Trade-in situation {mois_label(rv.per)}") if rv.trim else ("Traffic", mois_label(rv.per))
    s = slide_revue(ctx, rv, titre, sous_titre=sous, marque=marque, puces=puces[:5])
    g.texte(s, 0.25, 1.25, 2.0, 0.3, "Brand website", taille=14, gras=True, couleur="000000")
    g.texte(s, 2.5, 1.25, 3.0, 0.32, "Trade-in website", taille=14, gras=True, couleur="000000")
    g.texte(s, 2.5, 1.55, 4.0, 0.22, "Sum up (Period set above vs Y-1)", taille=8, italique=True, gras=True, couleur="595959")
    g.texte(s, 0.25, 1.55, 2.0, 0.22, "Brand website visits", taille=8, italique=True, gras=True, couleur="595959")
    mm, yy = (rv.vs_prec, "YoY")
    # visites du site de la marque tout à gauche, puis l'entonnoir du site de reprise : Sessions -> In Journey -> GA leads ; tuiles assez larges pour que « ▼ -30.5% QoQ » tienne sur une ligne
    LT = 2.0
    _carre(s, 0.25, 1.85, LT, 0.96, "Sessions", "n/a" if not site else nb(site), [(mm, e_site[0]), (yy, e_site[1])], 20)
    _carre(s, 2.5, 1.85, LT, 0.96, "Sessions", nb(ses), [(mm, e_s[0]), (yy, e_s[1])], 20)
    _carre(s, 5.0, 1.85, LT, 0.96, "In Journey", "n/a" if ij is None else nb(ij), [(mm, e_ij[0]), (yy, e_ij[1])], 20)
    _carre(s, 7.5, 1.85, LT, 0.96, "GA leads", "n/a" if hl is None else nb(hl), [(mm, e_hl[0]), (yy, e_hl[1])], 20)
    for x, lab in ((4.5, t1), (7.0, t2)):
        _ligne(s, x, 2.17, x + 0.5, 2.17, fleche=True)
        g.texte(s, x - 0.05, 1.93, 0.6, 0.22, "n/a" if lab is None else f"{lab:.1f} %".replace(".", ","), taille=8, couleur="000000", align=PP_ALIGN.CENTER)
    _ligne(s, 3.5, 2.81, 3.5, 3.13); _ligne(s, 3.5, 3.13, 8.5, 3.13); _ligne(s, 8.5, 3.13, 8.5, 2.87, fleche=True)
    g.texte(s, 5.5, 3.17, 1.0, 0.22, "n/a" if cvr is None else f"{cvr:.1f} %".replace(".", ","), taille=8, couleur="000000", align=PP_ALIGN.CENTER)
    # part des entrées sur la home de la reprise venant du site de la marque / directes (« Brand Websites » / « Trade In Entry » de Looker)
    orig = (ctx.d.origines_entree(rv.per, marque) if rv.trim else None) or _entree_origines(ctx, marque, rv.per)
    if orig:
        g.texte(s, 9.6, 1.25, 1.5, 0.3, "Trade-in entries by origin", taille=8, gras=True, italique=True, couleur="595959")
        g.anneau(s, 9.55, 1.5, 1.55, 1.7, ["Brand Websites", "Trade In Entry"], [orig[0], orig[1]], couleurs=["4472C4", "5BC0BE"], taille=7, legende_bas=True)

    # tableau des canaux + anneau
    if ct:
        g.texte(s, 0.5, 3.62, 3.5, 0.3, "Traffic landing on Trade In", taille=11, gras=True, couleur="000000")
        g.texte(s, 3.25, 3.67, 0.8, 0.25, rv.vs_prec, taille=10, italique=True, gras=True, couleur="000000", align=PP_ALIGN.CENTER)
        g.texte(s, 4.45, 3.67, 0.8, 0.25, "YoY", taille=10, italique=True, gras=True, couleur="000000", align=PP_ALIGN.CENTER)
        lignes = [["Channel group", "Sessions", "% Δ", "% Δ"]]
        for i, c in enumerate(ordre[:10], 1):
            em, ey = _d(can[c], can_p.get(c)) if can_p else None, _d(can[c], can_n.get(c)) if can_n else None
            lignes.append([f"{i}. {c}", nb(can[c]), _txt_evol(em), _txt_evol(ey)])
        g.tableau(s, 0.75, 3.95, 4.6, lignes, largeurs=[1.9, 0.9, 0.9, 0.9], taille=8, hauteur_ligne=0.25)
        top = ordre[:9]
        reste = sum(can[c] for c in ordre[9:])
        g.anneau(s, 5.6, 3.6, 5.4, 3.35, top + (["Other"] if reste else []), [can[c] for c in top] + ([reste] if reste else []), taille=9)
    return s


def ytd(ctx, rv, marque):
    d = ctx.d
    an = rv.per.annee
    tot = [d.leads_mensuels(a, marque) for a in (an - 1, an)]
    vn = [d.leads_mensuels(a, marque, "VN") for a in (an - 1, an)]
    if not (sum(tot[0]) or sum(tot[1])):
        return None
    n = rv.per.indice
    dj = d.dernier_jour_leads()
    if dj and dj.startswith(str(an)):
        n = min(n, int(dj[5:7]))
    i, ip = n - 1, n - 2
    puces = []
    ytd_c, ytd_p = sum(tot[1][:n]), sum(tot[0][:n])
    e = evol(ytd_c, ytd_p)
    if e is not None:
        puces.append(f"{nb(ytd_c)} leads year to date in {an}, {pct(e, 0, True)} vs {an - 1} ({nb(ytd_p)}).")
    mois = mois_label(rv.per).split()[0]
    parts = [f"{mois} total leads {pct(evol(tot[1][i], tot[0][i]), 0, True)} YoY"]
    if ip >= 0 and tot[1][ip]:
        parts.append(f"{pct(evol(tot[1][i], tot[1][ip]), 0, True)} vs {mois_label(rv.prec).split()[0]}")
    if vn[0][i]:
        parts.append(f"NC leads {pct(evol(vn[1][i], vn[0][i]), 0, True)} YoY")
    if tot[0][i]:
        puces.append(", ".join(parts) + ".")
    if ytd_c:
        puces.append(f"NC leads represent {sum(vn[1][:n]) / ytd_c * 100:.0f}% of {an} year-to-date leads ({sum(vn[0][:n]) / ytd_p * 100:.0f}% in {an - 1})." if ytd_p else
                     f"NC leads represent {sum(vn[1][:n]) / ytd_c * 100:.0f}% of {an} year-to-date leads.")
    s = slide_revue(ctx, rv, "YTD Leads Trends", marque=marque, puces=puces[:3])
    graphique_annuel(s, 0.4, 1.2, 10.4, 2.85, "Total leads year trend", tot[0], tot[1], n, an)
    graphique_annuel(s, 0.4, 4.05, 10.4, 2.9, "NC leads year trend", vn[0], vn[1], n, an)
    return s
