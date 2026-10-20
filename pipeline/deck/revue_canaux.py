"""Revue mensuelle France : « Leads by marketing channel » (matrice SOURCE_ACQ + colonnes empilées) et « Leads from main website » (matrice par emplacement de CTA).

Reproduit les diapositives 18/19, 24/25… du deck de référence avec des formes et graphiques natifs. Chaque fonction renvoie la diapositive, ou None sans données.
"""

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from . import graphiques as g
from .donnees import PROJET_BO, PROJETS, groupe_source

# groupes de sources du deck de référence (Tableau) : SOL et eDealer sont des sources à part (le regroupement standard de Donnees les met dans Display / Autre)
ORDRE_SOURCES = ["Display", "Edealer", "Emailing", "Main-Website", "Search", "Social", "SOL", "Autre"]      # (LLM est replié dans Autre : pas de ligne LLM dans le deck)
from .periodes import MOIS_COURTS
from .revue_commun import C_SANS, C_VN, C_VO, court, evol, mois_label, nb, pct, slide_revue

PROJETS = ["VN", "SANS", "VO"]
NOMS_PROJETS = {"VN": "New cars & Demo cars", "SANS": "Trade-in only", "VO": "Used cars"}
COUL = {"VN": C_VN, "SANS": C_SANS, "VO": C_VO}
BEIGE, GRIS_TOTAL = "FBF3EA", "E7E6E6"
NOM_SOURCE = {"Search": "SEA", "Autre": "Other"}


# ------------------------------------------------------------------ outils de dessin
def _rect(s, x, y, w, h, couleur):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb = g.rgb(couleur); r.line.fill.background(); r.shadow.inherit = False
    return r


def _t(s, x, y, w, h, txt, taille=8, gras=False, couleur="000000", align=PP_ALIGN.CENTER, italique=False):
    return g.texte(s, x, y, w, h, txt, taille=taille, gras=gras, couleur=couleur, align=align, italique=italique)


def legende_projets(s):
    """Légende « Purchase Project » en haut à droite (position relevée sur la diapositive 18 de référence)."""
    _rect(s, 8.85, 0.15, 1.35, 0.78, "F2F2F2")
    _t(s, 8.88, 0.16, 1.3, 0.2, "Purchase Project", taille=7.5, align=PP_ALIGN.LEFT, couleur="404040")
    for i, p in enumerate(PROJETS):
        y = 0.4 + 0.17 * i
        _rect(s, 8.9, y + 0.03, 0.1, 0.1, COUL[p])
        _t(s, 9.03, y - 0.02, 1.2, 0.2, NOMS_PROJETS[p] if p != "VN" else "New cars & Demo cars", taille=7, align=PP_ALIGN.LEFT, couleur="595959")


def matrice(s, x, y, w, h, lib_col0, lignes, colonnes, blocs, annees=None, hauteur_haut=0.5, taille=8, titre_col0="", hl_max=0.62):
    """Matrice « Tableau » : `lignes` = libellés (la dernière = Total), `colonnes` = [{bloc, valeurs[len(lignes)], mois, fond(bool), couleur}],
    `blocs` = {indice: titre}, `annees` = [(première_colonne, dernière_colonne, libellé)]. Chaque cellule : nombre, part de la colonne et barre colorée."""
    nl, nc = len(lignes), len(colonnes)
    wl = lib_col0
    wc = (w - wl) / nc
    h_bas = 0.22
    hl = min(hl_max, (h - hauteur_haut - h_bas) / nl)
    h = hauteur_haut + hl * nl + h_bas
    y0 = y + hauteur_haut
    for j, c in enumerate(colonnes):                                    # fonds beiges (année N) puis ligne de total
        if c["fond"]:
            _rect(s, x + wl + j * wc, y0, wc, hl * nl, BEIGE)
    _rect(s, x, y0 + hl * (nl - 1), w, hl, GRIS_TOTAL)
    for j, c in enumerate(colonnes):
        tot = c["valeurs"][-1] or 0
        vmax = max([v for v in c["bloc_vals"] if v] or [1])
        for i, v in enumerate(c["valeurs"]):
            if v is None:
                continue
            ytop = y0 + i * hl
            dernier = i == nl - 1
            _t(s, x + wl + j * wc, ytop + 0.02, wc, 0.2, nb(v), taille=taille + (1 if dernier else 0), gras=dernier, italique=dernier)
            part = f"{v / tot * 100:.0f}%" if tot else "–"
            _t(s, x + wl + j * wc, ytop + 0.2, wc, 0.2, part, taille=taille - 0.5, gras=dernier, italique=dernier, couleur="404040")
            bh = 0.14 if dernier else max(0.025, min(0.45, 0.2 * hl / 0.62) * (v / vmax if vmax else 0))      # la barre grandit avec la hauteur de ligne
            _rect(s, x + wl + j * wc + wc * 0.1, ytop + hl - bh - 0.03, wc * 0.8, bh, c["couleur"])
    for i, lab in enumerate(lignes):
        dernier = i == nl - 1
        _t(s, x, y0 + i * hl + hl / 2 - 0.12, wl, 0.25, "Total" if dernier else lab, taille=taille + (0.5 if dernier else 0), italique=dernier,
           align=PP_ALIGN.LEFT, couleur="000000" if dernier else "595959")
    _t(s, x, y + hauteur_haut - 0.27, wl, 0.22, titre_col0, taille=7, gras=True, align=PP_ALIGN.LEFT)
    # blocs et années
    for b, titre in blocs.items():
        idx = [j for j, c in enumerate(colonnes) if c["bloc"] == b]
        if not idx:
            continue
        _t(s, x + wl + idx[0] * wc, y + 0.05 if annees else y + 0.12, wc * len(idx), 0.25, titre, taille=9)
        if b > 0:
            ln = s.shapes.add_connector(1, Inches(x + wl + idx[0] * wc), Inches(y + 0.05), Inches(x + wl + idx[0] * wc), Inches(y + h))
            ln.line.color.rgb = g.rgb("D9D9D9"); ln.line.width = 12700
    for a, b2, lab in annees or []:
        _t(s, x + wl + a * wc, y + 0.27, wc * (b2 - a + 1), 0.2, lab, taille=7.5, couleur="595959")
    for j, c in enumerate(colonnes):
        _t(s, x + wl + j * wc, y + h - h_bas, wc, h_bas, c["mois"], taille=7.5, couleur="595959")
    return hl


# ------------------------------------------------------------------ données
def _mois_court(per):
    return court(per)


def _pcts(a, b):
    e = evol(a, b)
    return e


def _tot(d):
    return sum(d.values())


def groupe_revue(brut):
    b = (brut or "").strip().lower()
    g_ = "SOL" if b == "sol" else "Edealer" if b == "edealer" else groupe_source(brut)
    return "Autre" if g_ == "LLM" else g_


def _sources(ctx, per, marque):
    """{groupe: {projet: n}} des leads valides de la marque sur le mois, avec SOL et eDealer séparés (comme le deck de référence)."""
    try:
        out = {g: {p: 0 for p in PROJETS} for g in ORDRE_SOURCES}
        vu = False
        for r in ctx.d._periode_lignes("leads_acq", marque, per):
            if r["statut"] == "valide":
                out[groupe_revue(r["source_acq"])][PROJET_BO.get(r["projet_achat"], "SANS")] += int(r["n"])
                vu = True
        return out if vu else None
    except Exception:
        return None


# ------------------------------------------------------------------ 1. Leads by marketing channel
def canaux(ctx, rv, marque):
    """« Leads by marketing channel » : matrice source × projet × (N-1, N) × (M-1, M) + colonnes empilées par projet."""
    per_list = [rv.prec.n1(), rv.n1, rv.prec, rv.per]            # N-1 M-1, N-1 M, N M-1, N M
    data = [_sources(ctx, p, marque) for p in per_list]
    if all(dd is None for dd in data):
        return None
    vide = {g: {p: 0 for p in PROJETS} for g in ORDRE_SOURCES}
    data = [dd if dd is not None else vide for dd in data]          # marque lancée depuis : période de comparaison sans lead = 0
    tp = [{p: sum(dd[sg][p] for sg in dd) for p in PROJETS} for dd in data]
    if not sum(sum(t.values()) for t in tp):
        return None
    for p_, t_ in zip(per_list, tp):
        for k_, nom_ in (("VN", "leads_vn"), ("SANS", "leads_sans"), ("VO", "leads_vo")):
            ctx.faits.ajoute(nom_, t_[k_], marque, p_.etiquette("en"), "lead_deep_dive")
        ctx.faits.ajoute("leads", sum(t_.values()), marque, p_.etiquette("en"), "lead_deep_dive")
    srcs = [sg for sg in ORDRE_SOURCES if sg in data[0] and any(sum(dd[sg].values()) for dd in data)]
    lignes = [NOM_SOURCE.get(sg, sg) for sg in srcs] + ["Total"]
    colonnes = []
    for b, pj in enumerate(PROJETS):
        for k, dd in enumerate(data):
            vals = [dd[sg][pj] if dd[sg][pj] else None for sg in srcs] + [tp[k][pj]]
            colonnes.append({"bloc": b, "valeurs": vals, "mois": _mois_court(per_list[k]), "fond": k >= 2, "couleur": COUL[pj], "bloc_vals": vals[:-1]})
    titre, sous = ("LEAD DEEP DIVE", f"Trade-in situation {mois_label(rv.per)}") if rv.trim else ("Leads by marketing channel", mois_label(rv.per))
    s = slide_revue(ctx, rv, titre, sous_titre=sous, marque=marque, puces=_puces_canaux(data, tp, srcs, rv))
    legende_projets(s)
    annees = [(b * 4 + 2 * a, b * 4 + 2 * a + 1, str((per_list[0].annee, per_list[2].annee)[a])) for b in range(3) for a in (0, 1)]
    hl = matrice(s, 0.15, 1.3, 7.35, 5.65, 0.75, lignes, colonnes, {i: NOMS_PROJETS[p] for i, p in enumerate(PROJETS)}, annees=annees, titre_col0="SOURCE_ACQ", hl_max=1.0)
    # colonnes empilées à droite
    cats = [f"{_mois_court(p)}\n{nb(_tot(t))}" for p, t in zip(per_list, tp)]
    series, ordre = [], ["VO", "SANS", "VN"]
    haut = max(_tot(t) for t in tp) or 1
    for pj in ordre:
        vals = [t[pj] for t in tp]
        # pas d'étiquette sur un segment trop bas pour la porter (elle chevaucherait celle du segment voisin)
        etiq = [f"{nb(v)}\n{v / _tot(t) * 100:.0f}%" if v and _tot(t) and v / haut >= 0.08 else None for v, t in zip(vals, tp)]
        series.append((NOMS_PROJETS[pj], vals, COUL[pj], etiq))
    y_haut = 1.3 + 0.5                                              # haut de la première ligne de la matrice
    bas = 1.3 + 0.5 + hl * len(lignes) + 0.22                       # bas de la matrice (libellés des périodes) : les deux graphiques ont les mêmes bords haut et bas
    g.colonnes(s, 7.6, y_haut, 3.4, bas - y_haut, cats, [(n, v, c) for n, v, c, _ in series], empile=True, legende=False, taille=8,
               libelles_perso={i: e for i, (_, _, _, e) in enumerate(series)})
    _t(s, 7.9, 1.45, 1.5, 0.25, str(per_list[0].annee), taille=9, couleur="595959")
    _t(s, 9.4, 1.45, 1.5, 0.25, str(per_list[2].annee), taille=9, couleur="595959")
    return s


def _puces_canaux(data, tp, srcs, rv):
    """5 puces calculées : total vs M-1 et N-1, NC, sources principales NC, SEA (trade-in only) et Display."""
    out = []
    cur, prec, n1 = _tot(tp[3]), _tot(tp[2]), _tot(tp[1])
    em, ey = evol(cur, prec), evol(cur, n1)
    if em is not None:
        mot = lambda e: "increased" if e > 0 else "decreased"
        txt = f"Leads have {mot(em)} by {pct(em, 0, True)} {'Quarter' if rv.trim else 'Month'}-over-{'Quarter' if rv.trim else 'Month'}"
        if ey is not None:
            txt += f" and {mot(ey)} by {pct(ey, 0, True)} Year-over-Year"
        out.append(txt)
    nc, nc1 = tp[3]["VN"], tp[1]["VN"]
    ey = evol(nc, nc1)
    if nc:
        out.append(f"NC leads at {nc / 1000:.1f}k" + (f" ({pct(ey, 0, True)} YoY)" if ey is not None else "") if nc >= 1000 else
                   f"NC leads at {nb(nc)}" + (f" ({pct(ey, 0, True)} YoY)" if ey is not None else ""))
        top = sorted(srcs, key=lambda sg: -data[3][sg]["VN"])
        top = [t for t in top if data[3][t]["VN"]][:3]
        if top:
            out.append("Primary sources of NC leads: " + ", ".join(f"{NOM_SOURCE.get(t, t)} ({data[3][t]['VN'] / nc * 100:.0f}%)" for t in top))
    sw = data[3]["Search"]["SANS"] if "Search" in data[3] else 0
    if tp[3]["SANS"] and sw:
        sy = data[1]["Search"]["SANS"] / tp[1]["SANS"] * 100 if tp[1]["SANS"] else None
        out.append(f"SEA campaigns generate {sw / tp[3]['SANS'] * 100:.0f}% of trade-in only leads" + (f" (vs {sy:.0f}% LY)" if sy is not None else ""))
    dp, dl = _tot(data[3]["Display"]), _tot(data[1]["Display"])
    if dp or dl:
        e = evol(dp, dl)
        out.append("Display " + (f"{'up' if e > 0 else 'down'} {pct(e, 0, True)} vs LY" if e is not None else "no leads LY") + f" ({nb(dp)} leads)")
    return out[:5]


# ------------------------------------------------------------------ 2. Leads from main website
# codes de support (utm_content) du site de la marque, regroupés comme le « Support (groupe) » du deck de référence : plusieurs codes sont tronqués à la collecte
# (« er-menu » = header-menu, « hp-b » = hp-block, « hp-er-bar » = hp-header-bar) et les variantes « -vu » (visuel) se rattachent à leur emplacement
GROUPES_SUPPORT = [("pop", "Popin"), ("room", "showroom"), ("offer", "Offer-page"), ("footer", "Footer"), ("sticky", "Sticky bar"), ("menu", "Header menu"),
                   ("header-bar", "HP header bar"), ("er-bar", "HP header bar"), ("hp-b", "HP block"), ("slice", "Slice"), ("home", "Homepage"), ("slider", "Homepage"), ("cta", "CTA")]
ORDRE_SUPPORTS = ["showroom", "Popin", "Offer-page", "Footer", "Header menu", "CTA", "Homepage", "Slice", "Sticky bar", "HP block", "HP header bar"]   # ordre des lignes du deck


def groupe_support(brut):
    b = (brut or "").strip().lower()
    for motif, nom in GROUPES_SUPPORT:
        if motif in b:
            return nom
    return (brut or "").strip() or "(non renseigné)"


def _groupe_cta(cta):
    """{support brut: {projet: n}} -> {groupe de support: {projet: n}} (somme des codes d'un même emplacement)."""
    out = {}
    for brut, par_projet in cta.items():
        g_ = out.setdefault(groupe_support(brut), {p: 0 for p in par_projet})
        for p, n in par_projet.items():
            g_[p] = g_.get(p, 0) + n
    return out


def _leads_site_principal(ctx, per, marque):
    """{support brut: {projet: n}} de TOUS les leads valides venant du site principal de la marque (groupe de sources Main-Website, toutes campagnes)."""
    out = {}
    for r in ctx.d._periode_lignes("leads_acq", marque, per):
        if r["statut"] == "valide" and groupe_source(r["source_acq"]) == "Main-Website":
            sup = (r["support"] or "").strip() or "(non renseigné)"
            out.setdefault(sup, {p: 0 for p in PROJETS})[PROJET_BO.get(r["projet_achat"], "SANS")] += int(r["n"])
    return out


def site_principal(ctx, rv, marque):
    """« Leads from main website » : emplacements des CTA du site principal × projet × (M-1, M) de l'année N."""
    try:
        cur, prec = _groupe_cta(_leads_site_principal(ctx, rv.per, marque)), _groupe_cta(_leads_site_principal(ctx, rv.prec, marque))
    except Exception:
        return None
    if not cur and not prec:
        return None
    supports = sorted(set(cur) | set(prec), key=lambda k: (ORDRE_SUPPORTS.index(k) if k in ORDRE_SUPPORTS else 99, -(_tot(cur.get(k, {})) + _tot(prec.get(k, {})))))
    if not supports:
        return None
    z = {p: 0 for p in PROJETS}
    colonnes = []
    for b, pj in enumerate(PROJETS):
        for k, dd in enumerate((prec, cur)):
            vals = [dd.get(sp, z)[pj] or None for sp in supports]
            vals.append(sum(dd.get(sp, z)[pj] for sp in supports))
            colonnes.append({"bloc": b, "valeurs": vals, "mois": _mois_label_long(rv.prec if k == 0 else rv.per), "fond": False, "couleur": COUL[pj], "bloc_vals": vals[:-1]})
    s = slide_revue(ctx, rv, "Leads from main website", marque=marque, puces=_puces_site(cur, prec, supports))
    legende_projets(s)
    matrice(s, 0.15, 1.3, 10.8, 5.65, 1.3, supports + ["Total"], colonnes, {i: NOMS_PROJETS[p] for i, p in enumerate(PROJETS)},
            annees=[(0, 5, str(rv.per.annee))], hauteur_haut=0.55, titre_col0="Support (groupe)", hl_max=0.85)
    return s


def _mois_label_long(per):
    return per.etiquette("en").split()[0]


def _puces_site(cur, prec, supports):
    out = []
    vn = {sp: cur.get(sp, {}).get("VN", 0) for sp in supports}
    tot = sum(vn.values())
    if tot:
        top = [sp for sp in sorted(vn, key=lambda k: -vn[k]) if vn[sp]][:3]
        out.append(f"{top[0]} is the leading generator of NC leads ({vn[top[0]] / tot * 100:.0f}%)"
                   + (f", followed by {' and '.join(top[1:])}" if len(top) > 1 else ""))
        vp = {sp: prec.get(sp, {}).get("VN", 0) for sp in supports}
        tp_ = sum(vp.values())
        if tp_:
            ch = max(supports, key=lambda k: abs(vn[k] / tot - vp[k] / tp_))
            d = (vn[ch] / tot - vp[ch] / tp_) * 100
            if abs(d) >= 1:
                out.append(f"Largest share move vs previous month: {ch} ({d:+.0f} pts of NC leads)")
    return out
