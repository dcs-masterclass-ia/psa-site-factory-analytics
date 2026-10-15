"""Format « revue mensuelle France » : socle commun (contexte, pied de page, logo, drapeau, encadré de commentaires).

Reproduit le deck de référence « Stellantis FR — Trade in review » (73 diapositives, mois M) avec les données du mois choisi. Chaque module de diapositives
(revue_global.py, revue_marque.py, revue_canaux.py, revue_conquete.py) reçoit `(ctx, rv, …)` où `rv = Revue(ctx)` donne les périodes (mois M, M-1, M-12), les éléments
statiques (logos de marque, drapeau, repris du deck de référence et stockés dans le repo de données PRIVÉ : presentations/assets/revue_fr/) et les constructeurs de diapositive.
"""

import json
from datetime import date

from pptx.util import Inches, Pt

from . import graphiques as g
from .periodes import MOIS_COURTS

# ordre des marques dans le deck de référence (une section de 6 diapositives chacune) et groupes des diapositives « Traffic and leads by brand »
MARQUES = ["PEUGEOT", "CITROEN", "DS", "OPEL", "FIAT", "ALFA ROMEO", "JEEP", "LANCIA"]
GROUPE_XP = ["PEUGEOT", "CITROEN", "DS", "OPEL"]
GROUPE_XF = ["ALFA ROMEO", "FIAT", "FIAT PRO", "JEEP", "LANCIA"]
SECTIONS = ["EXECUTIVE SUMMARY", "TRAFFIC & LEAD DEEP DIVE", "PROJECTION CURRENT MONTH", "SALES PERFORMANCE", "PERFORMANCE BY BRAND"]

# couleurs du deck : barres 2025 (vert d'eau) / 2026 (orange), projets d'achat (bleu / rouge / orange), encadré de commentaires
C2025, C2026 = "76B7B2", "F28E2B"
C_VN, C_SANS, C_VO = "4E79A7", "E15759", "F28E2B"
BOITE_FOND, BOITE_BORD = "F3F8FF", "8FAADC"
MARINE = "043882"


def mois_label(per, langue="en"):
    """« August 2026 » (nom complet du mois, comme dans les titres du deck)."""
    from .periodes import MOIS_NOMS
    return f"{MOIS_NOMS[langue][per.indice - 1].capitalize()} {per.annee}"


class Revue:
    """Contexte du format : mois M (ctx.per), M-1, M-12, dossier des éléments statiques et date du deck."""

    def __init__(self, ctx):
        self.per = ctx.per
        if self.per.type != "mois":
            raise ValueError("Le format « revue mensuelle France » exige une période mensuelle.")
        self.prec = self.per.precedente()
        self.n1 = self.per.n1()
        self.dossier = ctx.d.hist.parent / "presentations" / "assets" / "revue_fr"
        try:
            self.manifeste = json.loads((self.dossier / "manifeste.json").read_text(encoding="utf-8"))
        except Exception:
            self.manifeste = {}
        j = ctx.aujourdhui
        self.date_deck = f"{MOIS_COURTS['en'][j.month - 1]} {j.year}"
        self.manquants = set()

    # -------------------------------------------------------------- éléments statiques
    def _image(self, slide, spec, **surcharge):
        if not spec:
            return None
        f = self.dossier / spec["fichier"]
        if not f.exists():
            self.manquants.add(spec["fichier"])
            return None
        x, y, w, h = (surcharge.get(k, spec[k]) for k in ("x", "y", "w", "h"))
        return slide.shapes.add_picture(str(f), Inches(x), Inches(y), Inches(w), Inches(h))

    def logo(self, slide, marque):
        """Logo de la marque, en haut au centre (position relevée dans le deck de référence)."""
        return self._image(slide, (self.manifeste.get("logos") or {}).get(marque))

    def drapeau(self, slide, grand=False):
        """Drapeau français à côté du logo autobiz (petit sur les diapositives de données, grand sur les séparateurs)."""
        return self._image(slide, self.manifeste.get("drapeau"), **({"x": 9.88, "y": 0.09, "w": 1.0, "h": 0.6} if grand else {}))


# ------------------------------------------------------------------ pied de page, titre, encadré
def pied(slide, rv, source="Source: autobiz"):
    """Pied de page du deck : date du deck à gauche, source à côté (10,5 pt)."""
    g.texte(slide, 0.14, 7.12, 1.0, 0.3, rv.date_deck, taille=9, couleur="000000")
    if source:
        g.texte(slide, 1.08, 7.15, 5.5, 0.28, source, taille=10.5, couleur="000000")


def boite_revue(slide, puces, x=11.12, y=1.15, w=2.11, h=5.78, taille=10):
    """Encadré de commentaires à droite des diapositives de données (fond F3F8FF, trait pointillé bleu clair, puces, texte centré verticalement)."""
    if not puces:
        return None
    return g.boite_pointillee(slide, x, y, w, h, puces, taille=taille, fond=BOITE_FOND, bord=BOITE_BORD, espace_apres=8, largeur_trait=1.0)


def slide_revue(ctx, rv, titre, sous_titre=None, marque=None, puces=None, source="Source: autobiz", taille_titre=24, notes=None, layout="Title and Subtitle", x_titre=0.14):
    """Nouvelle diapositive au format du deck de référence : titre (24 pt, police du gabarit), sous-titre facultatif en gras noir, logo de la marque au centre,
    drapeau, encadré de commentaires à droite (`puces`) et pied de page. Retourne la diapositive ; à elle d'ajouter les graphiques / tableaux
    dans la zone gauche (x de 0,15 à 11,0 po, y de 1,1 à 6,95 po)."""
    s = ctx.prs.slides.add_slide(ctx.lay[layout])
    ctx.n += 1
    s.shapes.title.text = titre
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(x_titre), Inches(0.27), Inches(8.0 if marque else 11.0), Inches(0.65)
    for para in t.text_frame.paragraphs:
        for r in para.runs:
            r.font.size = Pt(taille_titre)
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx == 12:
            if sous_titre:
                ph.text = sous_titre
                ph.left, ph.top, ph.width, ph.height = Inches(x_titre), Inches(0.62), Inches(8.0), Inches(0.4)
                for para in ph.text_frame.paragraphs:
                    for r in para.runs:
                        r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = g.rgb("000000")
            else:
                ph._element.getparent().remove(ph._element)
    if marque:
        rv.logo(s, marque)
    rv.drapeau(s)
    pied(s, rv, source)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    boite_revue(s, puces)
    return s


# ------------------------------------------------------------------ petits utilitaires de formatage (anglais, comme le deck)
def nb(v):
    return "–" if v is None else f"{int(round(v)):,}".replace(",", " ")


def evol(cur, ref):
    """Évolution en % entre deux valeurs (None si la base est nulle)."""
    return None if not ref else (cur - ref) / ref * 100


def pct(v, n=0, signe=False):
    if v is None:
        return "–"
    s = f"{abs(v):.{n}f}"
    return (("+" if v > 0 else "-" if v < 0 else "") if signe else "") + s + "%"


# ------------------------------------------------------------------ graphique commun : barres mensuelles N-1 / N (« Total leads year trend », « NC leads year trend »)
def graphique_annuel(slide, x, y, w, h, titre, an_prec, an_cur, n_ecoules, annee, taille=8):
    """Barres mensuelles de l'année précédente (vert d'eau) et de l'année en cours (orange) sur 12 mois ; `an_cur` n'est affiché que pour les `n_ecoules` premiers mois.
    L'étiquette de chaque barre de l'année en cours porte la valeur puis la croissance vs N-1 (« 34 660 / -1% »), comme dans les graphiques Tableau du deck.
    `an_prec`, `an_cur` : 12 valeurs (janvier..décembre). Un titre de graphique (gras, 10 pt) est posé au-dessus."""
    g.texte(slide, x, y, w, 0.25, titre, taille=10, gras=True, couleur="000000")
    mois = MOIS_COURTS["en"]
    cur = [v if i < n_ecoules else None for i, v in enumerate(an_cur)]
    etiq = []
    for i, v in enumerate(cur):
        if v is None:
            etiq.append(None)
            continue
        e = evol(v, an_prec[i])
        etiq.append(f"{nb(v)}\n{pct(e, 0, signe=True)}" if e is not None else nb(v))
    return g.colonnes(slide, x, y + 0.25, w, h - 0.25, mois, [(str(annee - 1), list(an_prec), C2025), (str(annee), cur, C2026)], taille=taille, libelles_perso={1: etiq},
                      legende=True, largeur_barre=40, format_nombre="#,##0")
