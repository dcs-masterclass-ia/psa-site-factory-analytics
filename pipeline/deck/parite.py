"""Contrôle « aucun chiffre perdu » entre deux versions d'une même présentation (ex. design V1 et design V2).

    python -m pipeline.deck.parite v1.pptx v2.pptx [--strict]

Pour chaque diapositive (même rang dans les deux fichiers), extrait tous les nombres affichés — textes, tableaux, valeurs des séries de graphiques natifs et texte alternatif
des graphiques en image (V2) — puis vérifie que chaque nombre de la version de référence se retrouve dans l'autre (arrondi toléré : 0,6 %, formats « 10.4k »).
Ne comparent que les nombres ≥ 100 ou suivis de « % » (les petits entiers isolés sont des puces, des pages, des jours). Code de sortie 1 avec --strict s'il manque un chiffre."""

import argparse
import re
import sys

from pptx import Presentation

JETON = re.compile(r"(?<![\w.])(\d{1,3}(?:[  ,]\d{3})+(?:\.\d+)?|\d+(?:[.,]\d+)?)(\s?[kK%])?")


def nombres(texte):
    """[(valeur, est_pourcentage)] des nombres d'un texte ; « 1 234 » et « 1,234 » = 1234 ; « 10.4k » = 10400."""
    out = []
    for m in JETON.finditer(texte or ""):
        t, suffixe = m.group(1), (m.group(2) or "").strip().lower()
        t = t.replace(" ", "").replace(" ", "")
        if re.fullmatch(r"\d{1,3}(,\d{3})+(\.\d+)?", t):
            t = t.replace(",", "")
        elif "," in t:
            t = t.replace(",", ".")
        try:
            v = float(t)
        except ValueError:
            continue
        if suffixe == "k":
            v *= 1000
        out.append((v, suffixe == "%"))
    return out


def _formes(formes):
    for sh in formes:
        yield sh
        if sh.shape_type == 6:
            yield from _formes(sh.shapes)


def extrait(slide):
    """Nombres affichés par une diapositive (hors pied de page)."""
    vus = []
    for sh in _formes(slide.shapes):
        if sh.top is not None and sh.top / 914400 >= 7.0:
            continue                                            # date, source, numéro de page
        if sh.has_text_frame:
            vus += nombres(sh.text_frame.text)
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    vus += nombres(c.text)
        if getattr(sh, "has_chart", False) and sh.has_chart:
            for pl in sh.chart.plots:
                for se in pl.series:
                    vus += [(float(v), False) for v in se.values if v is not None]
        if sh.shape_type == 13:
            vus += nombres(sh._element.nvPicPr.cNvPr.get("descr") or "")
    return vus


def present(v, est_pct, candidats):
    tol = max(0.55, 0.006 * abs(v))
    return any(abs(u - v) <= tol for u, _ in candidats)


def compare(chemin_ref, chemin_nouveau):
    a, b = Presentation(chemin_ref), Presentation(chemin_nouveau)
    manques, total = [], 0
    for i, (sa, sb) in enumerate(zip(a.slides, b.slides), 1):
        cible = extrait(sb)
        ref = [(v, p) for v, p in extrait(sa) if v >= 100 or p]
        ref = [(v, p) for v, p in ref if not (p is False and 2015 <= v <= 2035 and float(v).is_integer())]
        total += len(ref)
        absents = sorted({round(v, 1) for v, p in ref if not present(v, p, cible)})
        if absents:
            titre = (sa.shapes.title.text_frame.text if sa.shapes.title is not None else "").strip()[:40]
            manques.append((i, titre, absents))
    return len(a.slides), len(b.slides), total, manques


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("reference")
    ap.add_argument("nouveau")
    ap.add_argument("--strict", action="store_true")
    r = ap.parse_args()
    na, nb, total, manques = compare(r.reference, r.nouveau)
    print(f"{na} diapositives (référence) / {nb} (nouveau) ; {total} nombres contrôlés ; {sum(len(m[2]) for m in manques)} absents sur {len(manques)} diapositive(s)")
    for i, titre, absents in manques:
        print(f"  diapo {i} {titre!r} : {absents[:12]}{' …' if len(absents) > 12 else ''}")
    return 1 if (r.strict and (manques or na != nb)) else 0


if __name__ == "__main__":
    sys.exit(main())
