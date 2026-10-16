"""Format « revue mensuelle France » : résumé des résultats (diapositives 8, 9, 11 et 13 du deck de référence).

  - trafic_leads_groupe : « Executive Summary — Traffic and leads by brand xP / xF » (sessions et leads du mois, MoM et YoY) ;
  - projection_nc : « NC Leads projection » (leads New cars du mois en cours ramenés à la durée du mois) ;
  - conversion_ventes_nc : « NC Sales Conversion » (leads New cars et commandes NSC C1ST, cumul janvier -> mois M, N vs N-1).
Chaque fonction renvoie la diapositive ou None si aucune donnée ; une donnée absente s'affiche « n/a » ou une cellule vide, jamais un chiffre inventé.
"""

import calendar
from datetime import date, timedelta

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from . import graphiques as g
from .donnees import NOMS_MARQUES
from .periodes import MOIS_COURTS, MOIS_NOMS, Periode
from .revue_commun import GROUPE_XF, GROUPE_XP, evol, mois_label, nb, slide_revue

VERT_TXT, ROUGE_TXT = "00A651", "FF0000"
BLEU_PROJ = "4E79A7"
NA = "n/a"


# ------------------------------------------------------------------ utilitaires
def _nom(m):
    return NOMS_MARQUES.get(m, m.title())


def _sessions(d, per, m):
    """Sessions du site de reprise sur la période : GA4 direct, repli sur l'historique ; None si aucune donnée."""
    try:
        u = d.utilisateurs(per, m, [p for p in d.pays if p != "LU"] or d.pays)      # tableau « Traffic (Sessions) » des decks de référence : Belgique sans le Luxembourg (leads BELUX)
    except Exception:
        u = None
    if u and u.get("sessions"):
        return u["sessions"]
    try:
        v = sum(d.canaux(per, m).values())
    except Exception:
        v = 0
    return v or None


def _signe_pct(v, n=0):
    if v is None:
        return ""
    return f"{'+' if v > 0 else '-' if v < 0 else ''}{abs(v):.{n}f}%"


def _couleur(v):
    return VERT_TXT if (v or 0) > 0 else ROUGE_TXT if (v or 0) < 0 else "000000"


def _cellule(c, texte, taille=10, gras=False, couleur="000000", fond=None, align=PP_ALIGN.CENTER, sous=False):
    c.margin_left = c.margin_right = Inches(0.05)
    c.margin_top = c.margin_bottom = Inches(0.01)
    c.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = c.text_frame
    for i, ligne in enumerate(str(texte).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ligne
        r.font.size = Pt(taille); r.font.bold = gras; r.font.color.rgb = RGBColor.from_string(couleur)
    c.fill.solid()
    c.fill.fore_color.rgb = RGBColor.from_string(fond or "FFFFFF")


def _table(slide, x, y, largeurs, hauteurs, lignes):
    """Tableau natif ; lignes = liste de listes de dict(t, taille, gras, couleur, fond, align)."""
    gf = slide.shapes.add_table(len(lignes), len(largeurs), Inches(x), Inches(y), Inches(sum(largeurs)), Inches(sum(hauteurs)))
    gf.table.first_row = False
    gf.table.horz_banding = False
    for j, w in enumerate(largeurs):
        gf.table.columns[j].width = Inches(w)
    for i, h in enumerate(hauteurs):
        gf.table.rows[i].height = Inches(h)
    for i, row in enumerate(lignes):
        for j, spec in enumerate(row):
            spec = dict(spec)
            _cellule(gf.table.cell(i, j), spec.pop("t", ""), **spec)
    return gf


# ------------------------------------------------------------------ 8 et 9 : Traffic and leads by brand
def trafic_leads_groupe(ctx, rv, groupe, marques):
    d = ctx.d
    marques = [m for m in marques if m in d.marques]
    if not marques:
        return None
    per, prec, n1 = rv.per, rv.prec, rv.n1
    trafic = {m: (_sessions(d, per, m), _sessions(d, prec, m), _sessions(d, n1, m)) for m in marques}
    leads = {m: (d.leads_total(per, m), d.leads_total(prec, m), d.leads_total(n1, m)) for m in marques}
    if not any(v[0] for v in trafic.values()) and not any(v[0] for v in leads.values()):
        return None
    cour = MOIS_COURTS["en"]
    lab_cur = f"{cour[per.indice - 1]} {str(per.annee)[2:]}"
    lab_prec = f"{cour[prec.indice - 1]} {str(prec.annee)[2:]}"
    lab_n1 = f"{cour[n1.indice - 1]} {str(n1.annee)[2:]}"

    def ev(v):                                       # (cur, prec, n1) -> (MoM, YoY) en %
        cur, p, a = v
        return (evol(cur, p) if cur and p else None), (evol(cur, a) if cur and a else None)

    # commentaires calculés par règles
    puces = []
    ev_t = {m: ev(trafic[m]) for m in marques}
    ev_l = {m: ev(leads[m]) for m in marques}
    best = [(m, ev_t[m][1], ev_l[m][1]) for m in marques if ev_t[m][1] is not None and ev_l[m][1] is not None]
    if best:
        m, et, el = max(best, key=lambda x: x[1] + x[2])
        if et > 0 and el > 0:
            puces.append(f"{_nom(m)} stands out with strong YoY growth in both traffic ({_signe_pct(et)}) and leads ({_signe_pct(el)}).")
        worst = min(best, key=lambda x: x[1] + x[2])
        if worst[0] != m and worst[1] < 0 and worst[2] < 0:
            puces.append(f"{_nom(worst[0])} remains under pressure, with declines in both traffic ({_signe_pct(worst[1])} YoY) and leads ({_signe_pct(worst[2])} YoY).")
    mm_t_haut = [_nom(m) for m in marques if ev_t[m][0] is not None and ev_t[m][0] > 0 and ev_l[m][0] is not None and ev_l[m][0] < 0]
    if mm_t_haut:
        puces.append(f"{' and '.join(mm_t_haut)} show{'s' if len(mm_t_haut) == 1 else ''} positive MoM traffic trends but lower lead volumes in {MOIS_NOMS['en'][per.indice - 1].capitalize()}.")
    tot = [sum(leads[m][k] for m in marques) for k in range(3)]
    if len(puces) < 3 and tot[0] and tot[2]:
        puces.append(f"Total leads {groupe}: {nb(tot[0])} ({_signe_pct(evol(tot[0], tot[2]))} YoY).")
    puces = puces[:3]

    titre_sous = f"Traffic and leads by brand {groupe}"
    s = slide_revue(ctx, rv, "Executive Summary", sous_titre=titre_sous, puces=puces, source="Source: autobiz + GA", boite_pleine=True)
    s.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(28)

    n = len(marques)
    x0, x1 = 2.5, 10.5
    lw = (x1 - x0) / n

    def bloc(y, titre, sous, valeurs, dec):
        cadre = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.48), Inches(y), Inches(10.22), Inches(2.3))
        cadre.fill.solid(); cadre.fill.fore_color.rgb = g.rgb("FFFFFF")
        cadre.line.color.rgb = g.rgb("BFBFBF"); cadre.line.width = Pt(1); cadre.shadow.inherit = False
        paras = [(titre, {"taille": 24, "gras": True, "apres": 0, "couleur": "000000"})] + [(t, {"taille": 8, "gras": True, "apres": 0, "couleur": "000000"}) for t in sous]
        g.texte(s, 0.52, y + 0.03, 1.8, 0.9, paras)
        g.texte(s, 0.55, y + 1.28, 2.2, 0.3, f"MoM ({lab_cur} vs {lab_prec})", taille=11, gras=True, couleur="000000")
        g.texte(s, 0.55, y + 1.62, 2.2, 0.3, f"YoY ({lab_cur} vs {lab_n1})", taille=11, gras=True, couleur="000000")
        for i, m in enumerate(marques):
            cx = x0 + lw * i
            cur = valeurs[m][0]
            mom, yoy = ev(valeurs[m])
            g.texte(s, cx, y + 0.48, lw, 0.32, _nom(m), taille=14, couleur="404040", align=PP_ALIGN.CENTER)
            g.texte(s, cx, y + 0.8, lw, 0.45, nb(cur) if cur else "", taille=24, couleur="000000", align=PP_ALIGN.CENTER)
            for dy, v in ((1.28, mom), (1.62, yoy)):
                if v is not None:
                    fl = "▲ " if v > 0 else "▼ " if v < 0 else ""
                    g.texte(s, cx, y + dy, lw, 0.3, f"{fl}{_signe_pct(v, dec)}", taille=10, couleur=_couleur(v), align=PP_ALIGN.CENTER)

    bloc(1.6, "Traffic", ["(Sessions to", "Trade-in websites)"], trafic, 1)
    bloc(3.92, "Leads", ["(total leads)"], leads, 0)
    return s


# ------------------------------------------------------------------ 11 : projection du mois en cours
ORDRE_PROJECTION = ["CITROEN", "DS", "OPEL", "PEUGEOT", "ALFA ROMEO", "FIAT", "FIAT PRO", "JEEP", "LANCIA"]


def projection_nc(ctx, rv):
    d = ctx.d
    dernier = d.dernier_jour_leads()
    if not dernier:
        return None
    suiv = Periode("mois", rv.per.annee + (rv.per.indice == 12), 1 if rv.per.indice == 12 else rv.per.indice + 1)
    if not (suiv.debut.isoformat() <= dernier <= suiv.fin.isoformat()):
        return None                                         # pas de leads du mois en cours : rien à projeter
    jours_ecoules = int(dernier[8:10])
    jours_mois = calendar.monthrange(suiv.annee, suiv.indice)[1]

    class _P:
        debut, fin = suiv.debut, date.fromisoformat(dernier)
    n1 = suiv.n1()
    marques = [m for m in ORDRE_PROJECTION if m in d.marques]
    if not marques:
        return None
    # colonnes : M-1 (rv.prec), M (rv.per), projection du mois courant, même mois N-1, variation
    lignes_v, tot = [], [0, 0, 0.0, 0]
    for m in marques:
        a, b = d.leads(rv.prec, m)["VN"], d.leads(rv.per, m)["VN"]
        cur = d.leads(_P, m)["VN"]
        proj = cur / jours_ecoules * jours_mois
        r1 = d.leads(n1, m)["VN"]
        lignes_v.append((m, a, b, proj, r1))
        tot = [tot[0] + a, tot[1] + b, tot[2] + proj, tot[3] + r1]
    if not any(v[3] or v[1] or v[2] for v in lignes_v):
        return None
    mois_c = MOIS_NOMS["en"][suiv.indice - 1].capitalize()
    s = slide_revue(ctx, rv, mois_label(suiv), sous_titre="NC Leads projection", source=None,      # titre = le mois PROJETÉ (la référence affichait un titre périmé)
                    notes=f"Projection = leads New cars du 1er au {jours_ecoules} {mois_c} ÷ {jours_ecoules} jours × {jours_mois} jours (rythme moyen constant, indicatif).")
    # pied : « Source: autobiz data till <date> » avec la date surlignée en jaune
    tb = g.texte(s, 1.15, 7.12, 7.4, 0.3, "Source: autobiz data till ", taille=12, couleur="000000")
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = f"{mois_c} {jours_ecoules}{'th' if 11 <= jours_ecoules <= 13 else {1: 'st', 2: 'nd', 3: 'rd'}.get(jours_ecoules % 10, 'th')}"
    r.font.size = Pt(12); r.font.color.rgb = g.rgb("000000")
    rpr = r._r.get_or_add_rPr()
    hl = etree.SubElement(rpr, qn("a:highlight")); c = etree.SubElement(hl, qn("a:srgbClr")); c.set("val", "FFFF00")

    # en-têtes (années / mois) au-dessus du tableau
    x = 0.96
    lw = [1.7, 1.2, 1.2, 1.55, 1.2, 0.0]
    cols_x = {"label": x, "prec": x + 1.8, "per": x + 3.1, "proj": x + 4.5, "n1": x + 6.2, "var": 9.84}
    g.texte(s, cols_x["prec"], 1.2, 1.3, 0.3, str(rv.prec.annee), taille=11, couleur="404040", align=PP_ALIGN.CENTER)
    g.texte(s, cols_x["per"], 1.2, 1.3, 0.3, str(rv.per.annee), taille=11, couleur="404040", align=PP_ALIGN.CENTER)
    g.texte(s, cols_x["proj"], 1.2, 1.6, 0.3, str(suiv.annee), taille=11, gras=True, couleur=BLEU_PROJ, align=PP_ALIGN.CENTER)
    g.texte(s, cols_x["n1"], 1.2, 1.5, 0.3, str(n1.annee), taille=11, couleur="404040", align=PP_ALIGN.CENTER)
    g.texte(s, cols_x["prec"], 1.5, 1.3, 0.3, MOIS_NOMS["en"][rv.prec.indice - 1].lower(), taille=10, couleur="7F7F7F", align=PP_ALIGN.CENTER)
    g.texte(s, cols_x["per"], 1.5, 1.3, 0.3, MOIS_NOMS["en"][rv.per.indice - 1].lower(), taille=10, couleur="7F7F7F", align=PP_ALIGN.CENTER)
    g.texte(s, cols_x["n1"], 1.5, 1.5, 0.3, MOIS_NOMS["en"][n1.indice - 1].lower(), taille=10, couleur="7F7F7F", align=PP_ALIGN.CENTER)
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(cols_x["proj"] + 0.05), Inches(1.5), Inches(1.5), Inches(0.28))
    box.fill.solid(); box.fill.fore_color.rgb = g.rgb("FFFFFF"); box.line.color.rgb = g.rgb("000000"); box.line.width = Pt(1.5); box.shadow.inherit = False
    tf = box.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    pp = tf.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
    rr = pp.add_run(); rr.text = f"Projection {MOIS_COURTS['en'][suiv.indice - 1]}"; rr.font.size = Pt(12); rr.font.bold = True; rr.font.color.rgb = g.rgb("000000")
    vh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(cols_x["var"]), Inches(1.36), Inches(0.75), Inches(0.55))
    vh.fill.solid(); vh.fill.fore_color.rgb = g.rgb("4472C4"); vh.line.fill.background(); vh.shadow.inherit = False
    tf = vh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, t in enumerate(("Var", f"{str(n1.annee)}/{str(suiv.annee)[2:]}")):
        pp = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pp.alignment = PP_ALIGN.CENTER
        rr = pp.add_run(); rr.text = t; rr.font.size = Pt(9); rr.font.bold = True; rr.font.color.rgb = g.rgb("FFFFFF")

    def ligne(nom, a, b, proj, r1, gras=False, i=0):
        fond = "F2F2F2" if i % 2 else "FFFFFF"
        tc = "000000"
        return [{"t": nom, "taille": 10.5, "gras": not gras, "align": PP_ALIGN.LEFT, "fond": fond},
                {"t": nb(a), "taille": 10.5, "align": PP_ALIGN.RIGHT, "fond": fond, "couleur": tc},
                {"t": nb(b), "taille": 10.5, "align": PP_ALIGN.RIGHT, "fond": fond, "couleur": tc},
                {"t": nb(proj), "taille": 10.5, "align": PP_ALIGN.RIGHT, "fond": fond, "couleur": BLEU_PROJ, "gras": gras},
                {"t": nb(r1), "taille": 10.5, "align": PP_ALIGN.RIGHT, "fond": fond, "couleur": tc}]
    corps = [ligne(_nom_projection(m), a, b, pr, r1, i=i) for i, (m, a, b, pr, r1) in enumerate(lignes_v)]
    corps.append(ligne("Total général", tot[0], tot[1], tot[2], tot[3], gras=True, i=0))
    n = len(corps)
    h = min(0.43, 5.1 / n)
    _table(s, 0.96, 1.95, [1.8, 1.3, 1.4, 1.7, 1.7], [h] * n, corps)
    # colonne de variation
    vl = []
    for i, (m, a, b, pr, r1) in enumerate(lignes_v + [("T", tot[0], tot[1], tot[2], tot[3])]):
        e = evol(pr, r1)
        vl.append([{"t": _signe_pct(e) if e is not None else NA, "taille": 9, "gras": True, "couleur": _couleur(e) if e is not None else "7F7F7F",
                    "fond": "CFD5EA" if i % 2 == 0 else "E9EBF5"}])
    _table(s, cols_x["var"], 1.95, [0.75], [h] * n, vl)
    return s


def _nom_projection(m):
    return m


# ------------------------------------------------------------------ 13 : NC Sales Conversion
class _Cumul:
    """Mois de janvier à `fin` d'une année (interface de Periode.mois() pour Donnees.c1st_nsc)."""

    def __init__(self, annee, fin):
        self._m = [(annee, m) for m in range(1, fin + 1)]

    def mois(self):
        return self._m


def conversion_ventes_nc(ctx, rv):
    d = ctx.d
    per = rv.per
    fin_m = per.indice
    a, a1 = per.annee, per.annee - 1
    jours_fin = calendar.monthrange(a, fin_m)[1]
    ordre = [("xP", ["CITROEN", "PEUGEOT", "DS", "OPEL"]), ("Xf", ["ALFA ROMEO", "FIAT PRO", "FIAT", "JEEP", "LANCIA"])]
    infos = {}
    for _, ms in ordre:
        for m in ms:
            if m not in d.marques:
                continue
            # « # Leads Generated » = leads VN ENVOYÉS (Lead sent = YES), ceux que C1ST affiche ; repli sur tous les leads VN si l'historique des envoyés manque
            e26, e25 = d.leads_envoyes_mensuels(a, m, "VN"), d.leads_envoyes_mensuels(a1, m, "VN")
            l26 = sum((e26 if e26 is not None else d.leads_mensuels(a, m, "VN"))[:fin_m])
            l25 = sum((e25 if e25 is not None else d.leads_mensuels(a1, m, "VN"))[:fin_m])
            o26 = d.c1st_nsc(_Cumul(a, fin_m), m)
            o25 = d.c1st_nsc(_Cumul(a1, fin_m), m)
            infos[m] = (l26, l25, o26, o25)
    if not infos or not any(v[0] or v[1] for v in infos.values()):
        return None

    def agrege(ms):
        v = [infos[m] for m in ms if m in infos]
        if not v:
            return None
        l26, l25 = sum(x[0] for x in v), sum(x[1] for x in v)
        ok = all(x[2] is not None and x[3] is not None for x in v)
        return (l26, l25, sum(x[2] for x in v) if ok else None, sum(x[3] for x in v) if ok else None)
    groupes = {nom: agrege(ms) for nom, ms in ordre}
    tous = agrege([m for _, ms in ordre for m in ms])

    def tx(o, l):
        return o / l * 100 if (o is not None and l) else None

    def ligne(nom, v, sous=False, fond_tot=False):
        l26, l25, o26, o25 = v
        el = evol(l26, l25)
        t26, t25 = tx(o26, l26), tx(o25, l25)
        dt = t26 - t25 if t26 is not None and t25 is not None else None
        do = o26 - o25 if o26 is not None and o25 is not None else None
        fond = "D9D9D9" if sous else "FFFFFF"
        fn = lambda t, **k: dict({"t": t, "taille": 10, "gras": sous, "fond": fond, "align": PP_ALIGN.RIGHT}, **k)
        pc = lambda x: f"{x:.1f}%" if x is not None else NA
        return [fn(nom, align=PP_ALIGN.LEFT, gras=True if sous else False),
                fn(nb(l26)), fn(nb(l25)),
                fn(f"{el:.0f}%" if el is not None else NA, couleur=VERT_TXT if (el or 0) >= 0 else ROUGE_TXT, gras=True),
                fn(pc(t26)), fn(pc(t25)), fn(f"{dt:.1f}%" if dt is not None else NA),
                fn(nb(o26) if o26 is not None else NA, fond="FFD966" if fond_tot else fond, gras=sous), fn(nb(o25) if o25 is not None else NA, gras=sous),
                fn(nb(do) if do is not None else NA, fond="E2F0D9" if not sous else "E2F0D9", gras=sous)]

    entete = ["ALL", "# Leads Generated\n2026", "# Leads Generated\n2025", "Δ\nLeads Evol", "% NSC\nOrders\n2026", "% NSC\nOrders\n2025", "Δ\n% NSC\norders",
              "# NSC\nOrders\n2026", "# NSC\nOrders\n2025", "Δ\n#NSC\nOrders"]
    corps = [[{"t": h, "taille": 10, "gras": True, "couleur": "FFFFFF", "fond": "002060"} for h in entete]]
    for nom, ms in ordre:
        for m in ms:
            if m in infos:
                corps.append(ligne(_nom(m), infos[m]))
        if groupes[nom]:
            corps.append(ligne(nom, groupes[nom], sous=True))
    corps[0][0]["align"] = PP_ALIGN.LEFT
    if tous:
        corps.append(ligne("Total " + "/".join(rv.pays), tous, sous=True, fond_tot=True))
    largeurs = [1.4, 1.15, 1.15, 1.2, 0.85, 0.85, 0.85, 0.85, 0.85, 0.85]
    hauteurs = [0.62] + [0.255] * (len(corps) - 1)
    s = slide_revue(ctx, rv, "NC Sales Conversion", sous_titre=f"Leads from Jan. to {MOIS_NOMS['en'][fin_m - 1].capitalize()} {jours_fin}th. {str(a)[2:]} incl.", source="Source: autobiz")
    _table(s, 1.5, 1.2, largeurs, hauteurs, corps)

    # encadré de synthèse
    seuil = ctx.aujourdhui - timedelta(days=90)
    provisoire = per.fin > seuil
    xp, xf = groupes["xP"], groupes["Xf"]
    mois_txt = MOIS_COURTS["en"][fin_m - 1]
    lignes = [f"From Jan to {mois_txt} {jours_fin}th. {a}:"]
    parts = []
    if xp and evol(xp[0], xp[1]) is not None:
        parts.append(f"a {abs(evol(xp[0], xp[1])):.0f}% {'increase' if xp[0] >= xp[1] else 'decrease'} in NC leads for xP dealers")
    if xf and evol(xf[0], xf[1]) is not None:
        parts.append(f"a {abs(evol(xf[0], xf[1])):.0f}% {'increase' if xf[0] >= xf[1] else 'decrease'} for the xF dealers")
    l1 = "A" + (parts[0][1:] if parts else "") if parts else ""
    if parts:
        l1 = parts[0][0].upper() + parts[0][1:] + (" and " + parts[1] if len(parts) > 1 else "")
    if tous and tous[2] is not None:
        l1 = (l1 + " leading to " if l1 else "") + f"{nb(tous[2])} NC sales so far"
        lignes.append(l1 + (" (Sales are still processing)" if provisoire else ""))
    elif l1:
        lignes.append(l1 + " (NSC orders: n/a)")
    if tous and tous[2] is not None and tous[3] is not None:
        dd = tous[2] - tous[3]
        lignes.append(f"Already {nb(dd)} additional sales compared to last year" if dd > 0 else f"{nb(-dd)} fewer sales than last year" if dd < 0 else "Sales in line with last year")
    contrib = [(m, v[2] - v[3]) for m, v in infos.items() if v[2] is not None and v[3] is not None]
    if contrib:
        top = [m for m, dv in sorted(contrib, key=lambda x: -x[1]) if dv > 0][:3]
        if top:
            lignes.append(", ".join(_nom(m) for m in top[:-1]) + (" and " if len(top) > 1 else "") + _nom(top[-1]) + (" are the top performers" if len(top) > 1 else " is the top performer"))
    boite = g.boite_pointillee(s, 0.26, 5.43, 12.31, 1.57, lignes, taille=12, fond="F3F8FF", bord="9DC3E6", espace_apres=6, largeur_trait=1.0)
    p0 = boite.text_frame.paragraphs[0]._p
    ppr = p0.get_or_add_pPr()
    ppr.set("marL", "0"); ppr.set("indent", "0")
    for tag in ("a:buFont", "a:buChar"):
        for e in ppr.findall(qn(tag)):
            ppr.remove(e)
    etree.SubElement(ppr, qn("a:buNone"))
    if provisoire:                                       # « Sales are still processing » en gras souligné
        for para in boite.text_frame.paragraphs[1:]:
            run = para.runs[0] if para.runs else None
            if run is not None and "(Sales are still processing)" in run.text:
                avant, _, _ = run.text.partition("(Sales are still processing)")
                run.text = avant + "("
                r2 = para.add_run(); r2.text = "Sales are still processing"
                r3 = para.add_run(); r3.text = ")"
                for rr, extra in ((r2, True), (r3, False)):
                    rr.font.size = Pt(12); rr.font.name = "Calibri"; rr.font.color.rgb = g.rgb("000000")
                    if extra:
                        rr.font.bold = True; rr.font.underline = True
                break
    return s
