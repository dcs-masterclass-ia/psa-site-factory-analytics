"""« NC Sales Conversion » : vue tous marchés (retour DKAM du 20/10/2026), deux diapositives placées avant le tableau du pays :
  1. NC Leads YoY, répartition par marché (barres, part en %), carte d'Europe colorée par croissance vs N-1, commentaire calculé ;
  2. entonnoirs « # Leads Generated » contre « # NSC Orders » puis contre « % NSC Orders » par marché.
Données : c1st/kpi.json du repo de données (leads C1ST type 1, activité VN, campagne Autobiz ; Spoticar exclu, ses leads sont des rachats ; commandes = commandes NSC rattachées
au lead entre 1 et 90 jours ; taux de commande = commandes NSC ÷ clients uniques). Le Luxembourg est compté dans la Belgique côté C1ST (« Belux »). Fond de carte : europe.json (Natural Earth simplifié)."""

import calendar
import json
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.util import Emu, Inches, Pt

from . import graphiques as g
from .revue_commun import slide_revue

MARCHES = [("FR", "France", "FRANCE"), ("IT", "Italy", "ITALY"), ("ES", "Spain", "SPAIN"), ("BE", "Belux", "BELUX"), ("DE", "Germany", "GERMANY"),
           ("PT", "Portugal", "PORTUGAL"), ("PL", "Poland", "POLAND"), ("AT", "Austria", "AUSTRIA"), ("GB", "UK", "UK")]
ISO_MARCHE = {"FRA": "FR", "ITA": "IT", "ESP": "ES", "BEL": "BE", "LUX": "BE", "DEU": "DE", "PRT": "PT", "POL": "PL", "AUT": "AT", "GBR": "GB"}
ANCRES = {"FR": (2.4, 46.6), "IT": (12.2, 42.5), "ES": (-3.6, 40.0), "BE": (3.6, 53.4), "DE": (10.2, 51.2), "PT": (-10.2, 39.4), "PL": (19.2, 52.0), "AT": (14.2, 47.6), "GB": (-2.4, 53.4)}
NOM_LONG = {"BE": "Belgium"}
ORDINAL = {1: "first", 2: "second", 3: "third", 4: "fourth", 5: "fifth", 6: "sixth", 7: "seventh", 8: "eighth", 9: "ninth"}
CLASSES = [(-1e9, 0, "F28B7A"), (0, 50, "A8D5A0"), (50, 150, "6FB06F"), (150, 250, "4E9B5B"), (250, 1e9, "2E7D45")]
GRIS_FOND = "E3E6EA"
BLEU, ORANGE = "56B4DA", "E2622F"


def _couleur(croiss):
    for lo, hi, c in CLASSES:
        if lo <= croiss < hi:
            return c
    return GRIS_FOND


def donnees(ctx, rv):
    """[{code, nom, leads, leads_n1, commandes, uniques, taux}] triés par leads décroissants ; None si le relevé C1ST est absent ou incomplet."""
    brut = ctx.d._c1st()
    if not brut:
        return None
    fin = rv.per.mois()[-1]                         # (année, mois) de fin de période
    mois = lambda a: [f"{a}-{m:02d}" for m in range(1, fin[1] + 1)]
    partiel = getattr(ctx.d, "_c1st_partiel", "")
    out = []
    for code, nom, _ in MARCHES:
        d = brut.get(code) or {}
        def somme(a, section, champ):
            tot = 0
            for m in mois(a):
                e = d.get(m)
                if not e or m == partiel or e.get(section) is None:
                    return None
                tot += sum((r.get(champ) or 0) for r in e[section] if r.get("marque") != "SP")
            return tot
        l26, l25 = somme(fin[0], "leads", "leads"), somme(fin[0] - 1, "leads", "leads")
        o26, u26 = somme(fin[0], "nsc", "commandes_nsc"), somme(fin[0], "commandes", "clients_uniques")
        if l26 is None or l25 is None:
            continue
        out.append({"code": code, "nom": nom, "leads": l26, "leads_n1": l25, "commandes": o26, "uniques": u26,
                    "taux": (o26 / u26 * 100) if (o26 is not None and u26) else None})
    out = [x for x in out if x["leads"] or x["leads_n1"]]
    return sorted(out, key=lambda x: -x["leads"]) or None


def _fmt_k(v):
    return f"{v / 1000:.0f}K" if v >= 10000 else (f"{v / 1000:.1f}K".replace(".0K", "K") if v >= 1000 else f"{v:,}")


def _sous_titre(ctx, rv):
    a, m = rv.per.mois()[-1]
    jour = calendar.monthrange(a, m)[1]
    mois = ["Jan", "Feb", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"][m - 1]
    suff = "th" if 11 <= jour <= 13 else {1: "st", 2: "nd", 3: "rd"}.get(jour % 10, "th")
    return f"Leads from Jan. to {mois} {jour}{suff}. {str(a)[2:]} incl."


def _pied(ctx):
    d = ctx.d.c1st_date()
    return f"Source: autobiz — C1ST data as of {d}; NSC orders are linked to leads up to 90 days after creation: recent figures are provisional" if d else "Source: autobiz — C1ST"


# ------------------------------------------------------------------ diapo 1 : leads par marché + carte
def _carte(s, x0, y0, w, h, par_code):
    geo = json.loads((Path(__file__).resolve().parent / "europe.json").read_text())
    lon0, lon1, lat0, lat1 = -10.5, 24.5, 35.5, 59.0
    cosl = 0.669
    ech = min(w / ((lon1 - lon0) * cosl), h / (lat1 - lat0))
    ox = x0 + (w - (lon1 - lon0) * cosl * ech) / 2
    oy = y0 + (h - (lat1 - lat0) * ech) / 2
    pt = lambda lon, lat: (ox + (lon - lon0) * cosl * ech, oy + (lat1 - lat) * ech)
    E = lambda v: int(round(v * 914400))
    for iso, polys in geo.items():
        marche = ISO_MARCHE.get(iso)
        cible = par_code.get(marche) if marche else None
        fond = _couleur(cible["croiss"]) if cible and cible["croiss"] is not None else GRIS_FOND
        for poly in polys:
            pts = [pt(lo, la) for lo, la in poly]
            fb = s.shapes.build_freeform(E(pts[0][0]), E(pts[0][1]))
            fb.add_line_segments([(E(px), E(py)) for px, py in pts[1:]], close=True)
            sh = fb.convert_to_shape()
            sh.fill.solid(); sh.fill.fore_color.rgb = g.rgb(fond)
            sh.line.color.rgb = g.rgb("FFFFFF"); sh.line.width = Pt(0.5); sh.shadow.inherit = False
    for code, mk in par_code.items():                      # étiquettes : nom, volume, croissance
        lon, lat = ANCRES[code]
        px, py = pt(lon, lat)
        cr = f"{mk['croiss']:+.0f}%" if mk["croiss"] is not None else "n/a"
        g.texte(s, px - 0.6, py - 0.22, 1.2, 0.5, [(mk["etiq"], {"taille": 6.5, "gras": False, "couleur": "404040", "apres": 0}),
                                                   (f"{mk['leads']:,}", {"taille": 6.5, "couleur": "404040", "apres": 0}), (cr, {"taille": 6.5, "couleur": "404040", "apres": 0})],
                taille=6.5, align=PP_ALIGN.CENTER)
    return pt


def nc_marches(ctx, rv):
    marches = donnees(ctx, rv)
    if not marches:
        return None
    tot26, tot25 = sum(m["leads"] for m in marches), sum(m["leads_n1"] for m in marches)
    for m in marches:
        m["part"] = m["leads"] / tot26 * 100 if tot26 else 0
        m["croiss"] = (m["leads"] / m["leads_n1"] - 1) * 100 if m["leads_n1"] else None
        m["etiq"] = dict((c, e) for c, _, e in MARCHES)[m["code"]]
    par_code = {m["code"]: m for m in marches}
    ici = ctx.d.pays[0] if ctx.d.pays else None
    ici = {"LU": "BE"}.get(ici, ici)
    puces = []
    rang = next((i for i, m in enumerate(marches, 1) if m["code"] == ici), None)
    if rang:
        m = marches[rang - 1]
        puces.append(f"{m['nom']} ranks {ORDINAL.get(rang, str(rang) + 'th')} in NC lead generation, accounting for {m['part']:.0f}% of total autobiz leads.")
        if m["croiss"] is not None:
            puces.append(f"Year-to-date lead volume in {rv.per.annee} {'rose' if m['croiss'] >= 0 else 'fell'} {abs(m['croiss']):.0f}% compared with the same period last year.")
    baisse = [m["nom"] for m in marches if m["croiss"] is not None and m["croiss"] < 0]
    puces.append("NC lead volumes increased in every market" + (f" except {', '.join(baisse[:-1]) + (' and ' if len(baisse) > 1 else '') + baisse[-1]}." if baisse else "."))
    s = slide_revue(ctx, rv, "NC Sales Conversion", sous_titre=_sous_titre(ctx, rv), source=None, puces=puces)
    _t = lambda x, y, w, h, t, **k: g.texte(s, x, y, w, h, t, **k)
    # chiffre clé
    _t(0.5, 1.3, 3.2, 0.3, "NC Leads YoY", taille=11, gras=True, couleur="1F2A6B", align=PP_ALIGN.CENTER)
    e = (tot26 / tot25 - 1) * 100 if tot25 else None
    _t(0.5, 1.7, 3.2, 0.45, f"{tot26:,}", taille=20, couleur="2E7D32", align=PP_ALIGN.CENTER)
    if e is not None:
        _t(0.5, 2.15, 3.2, 0.4, f"{e:+.0f}%", taille=16, couleur="3FAE4A" if e >= 0 else "C0392B", align=PP_ALIGN.CENTER)
    # barres par marché
    _t(0.3, 2.85, 3.6, 0.3, "NC Leads by Country and Share %", taille=10, gras=True, couleur="1F2A6B", align=PP_ALIGN.CENTER)
    y = 3.25
    haut = max(m["leads"] for m in marches) or 1
    pas = min(0.38, 3.4 / len(marches))
    for m in marches:
        _t(0.3, y, 0.95, pas, m["etiq"], taille=7, couleur="404040", align=PP_ALIGN.RIGHT, ancre=MSO_ANCHOR.MIDDLE)
        largeur = max(0.04, 2.2 * m["leads"] / haut)
        b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(y + 0.05), Inches(largeur), Inches(pas - 0.1))
        b.fill.solid(); b.fill.fore_color.rgb = g.rgb("9DC9E8"); b.line.fill.background(); b.shadow.inherit = False
        tb = s.shapes.add_textbox(Inches(1.3 + largeur + 0.05), Inches(y), Inches(1.9), Inches(pas))       # volume puis part (en gras) à la suite de la barre
        tb.text_frame.margin_left = tb.text_frame.margin_right = 0
        tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        pp = tb.text_frame.paragraphs[0]
        r1 = pp.add_run(); r1.text = f"{m['leads']:,} / "; r1.font.size = Pt(8); r1.font.color.rgb = g.rgb("000000")
        r2 = pp.add_run(); r2.text = f"{m['part']:.0f}%"; r2.font.size = Pt(8); r2.font.bold = True; r2.font.color.rgb = g.rgb("000000")
        y += pas
    # carte
    _t(4.2, 1.3, 6.5, 0.3, "NC Leads by Country and YoY Growth", taille=11, gras=True, couleur="1F2A6B", align=PP_ALIGN.CENTER)
    _carte(s, 4.2, 1.6, 6.5, 5.25, par_code)
    crs = [m["croiss"] for m in marches if m["croiss"] is not None]
    if crs:                                                                 # légende « % growth vs Y-1 »
        lx, ly = 9.0, 6.55
        _t(lx - 0.3, ly - 0.25, 2.2, 0.2, "% growth vs Y-1", taille=6.5, gras=True, couleur="404040")
        for i, (_, _, c) in enumerate(CLASSES):
            r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(lx + i * 0.32), Inches(ly), Inches(0.32), Inches(0.14))
            r.fill.solid(); r.fill.fore_color.rgb = g.rgb(c); r.line.fill.background(); r.shadow.inherit = False
        _t(lx - 0.35, ly + 0.15, 0.7, 0.2, f"{min(crs):+.0f}%", taille=6.5, couleur="404040")
        _t(lx + 1.2, ly + 0.15, 0.7, 0.2, f"{max(crs):+.0f}%", taille=6.5, couleur="404040", align=PP_ALIGN.RIGHT)
    _t(1.08, 7.12, 11.0, 0.28, _pied(ctx), taille=8, couleur="000000")
    return s


# ------------------------------------------------------------------ diapo 2 : entonnoirs leads / commandes NSC / taux
def _entonnoir(s, x, y, w, h, noms, gauche, droite, fmt_g, fmt_d, titre_g, titre_d, mil=1.5):
    """Deux barres horizontales opposées (leads à gauche, commandes ou taux à droite), les noms de marché au centre."""
    lw = (w - mil) / 2
    for xx, vals, fmt, coul, rev in ((x, gauche, fmt_g, BLEU, True), (x + lw + mil, droite, fmt_d, ORANGE, False)):
        cd = CategoryChartData()
        cd.categories = noms
        cd.add_series("v", [v or 0 for v in vals])
        gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(xx), Inches(y), Inches(lw), Inches(h), cd)
        ch = gf.chart
        ch.has_legend = False; ch.has_title = False
        ch.font.size = Pt(8)
        ch.value_axis.visible = False; ch.value_axis.has_major_gridlines = False
        ch.value_axis.minimum_scale = 0
        ch.value_axis.maximum_scale = (max(v or 0 for v in vals) or 1) * 1.02
        ch.value_axis.reverse_order = rev                        # barres qui partent du centre vers la gauche
        ch.category_axis.visible = False
        ch.category_axis.reverse_order = True                    # premier marché en haut
        pl = ch.plots[0]
        pl.gap_width = 22
        ser = pl.series[0]
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = g.rgb(coul)
        haut = max(v or 0 for v in vals) or 1
        for i, v in enumerate(vals):
            dl = ser.points[i].data_label
            dl.position = XL_LABEL_POSITION.INSIDE_END if (v or 0) / haut > 0.25 else XL_LABEL_POSITION.OUTSIDE_END
            dl.text_frame.text = fmt(v) if v is not None else "n/a"
            for p in dl.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8); r.font.color.rgb = g.rgb("000000")
        g.zone_fixe(gf, 0.01, 0.0, 0.98, 1.0)
    # noms de marché au centre
    pas = h / len(noms)
    for i, n in enumerate(noms):
        g.texte(s, x + lw, y + i * pas, mil, pas, n, taille=8, couleur="000000", align=PP_ALIGN.CENTER, ancre=MSO_ANCHOR.MIDDLE)
    for xx, t in ((x, titre_g), (x + lw + mil, titre_d)):          # en-têtes façon liste déroulante
        r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(xx), Inches(y - 0.38), Inches(lw), Inches(0.28))
        r.fill.solid(); r.fill.fore_color.rgb = g.rgb("FFFFFF"); r.line.color.rgb = g.rgb("B9C2CC"); r.line.width = Pt(0.75); r.shadow.inherit = False
        tf = r.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
        run = p.add_run(); run.text = t; run.font.size = Pt(8); run.font.color.rgb = g.rgb("404040")


def nc_entonnoirs(ctx, rv):
    marches = donnees(ctx, rv)
    if not marches or not any(m["commandes"] for m in marches):
        return None
    noms = [{"UK": "United Kingdom", "Belux": "Belgium"}.get(m["nom"], m["nom"]) for m in marches]
    s = slide_revue(ctx, rv, "NC Sales Conversion", sous_titre=_sous_titre(ctx, rv), source=None)
    n = len(marches)
    h = 2.3
    _entonnoir(s, 0.35, 1.75, 12.4, h, noms, [m["leads"] for m in marches], [m["commandes"] for m in marches], _fmt_k, lambda v: _fmt_k(v),
               "# Leads Generated", "# NSC Orders")
    _entonnoir(s, 0.35, 4.75, 12.4, h, noms, [m["leads"] for m in marches], [m["taux"] for m in marches], _fmt_k,
               lambda v: f"{v:.1f}%".replace(".0%", "%"), "# Leads Generated", "% NSC Orders")
    _t = g.texte
    _t(s, 1.08, 7.12, 11.0, 0.28, _pied(ctx), taille=8, couleur="000000")
    return s
