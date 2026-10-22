"""Revue mensuelle France : diapositives « Executive Summary » (résumé exécutif, vue globale, tendance des leads toutes marques).

Reproduit les diapositives 4 à 6 du deck de référence avec les chiffres du mois choisi. Tout est calculé par règles à partir de `Donnees` ; une donnée absente
donne « n/a » ou une puce omise, jamais un chiffre inventé. Chaque fonction renvoie la diapositive, ou None si aucune donnée de leads.
"""

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

from . import graphiques as g
from .periodes import MOIS_COURTS, MOIS_NOMS
from .revue_commun import C_SANS, C_VN, C_VO, MARQUES, graphique_annuel, evol, mois_label, nb, pct, slide_revue

NAVY_TABLEAU = "000080"
GRIS_BANDEAU = "E6E6E6"
# couleurs de marque du tableau de bord (palette Tableau, comme la capture de référence)
COULEUR_MARQUE = {"PEUGEOT": "76B7B2", "SPOTICAR": "59A14F", "CITROEN": "4E79A7", "FIAT": "EDC948", "OPEL": "E15759", "JEEP": "FF9DA7", "DS": "F28E2B",
                  "STELLANTIS &YOU": "76B7B2", "ALFA ROMEO": "B07AA1", "LANCIA": "9C755F", "FIAT PRO": "BAB0AC"}


# ------------------------------------------------------------------ calculs communs
def _mois_nom(per):
    return MOIS_NOMS["en"][per.indice - 1].capitalize()


def _k(v):
    """54 917 -> « 54.9k », 535 271 -> « 535k », sous 1 000 : entier."""
    if v is None:
        return "n/a"
    if abs(v) >= 100000:
        return f"{v / 1000:.0f}k"
    if abs(v) >= 1000:
        return f"{v / 1000:.1f}k"
    return f"{int(round(v))}"


def _faits(ctx, rv):
    """Chiffres de base : mois, mois précédent, N-1 du mois, cumul de l'année et du même cumul N-1, parts New cars. None sans aucun lead."""
    d, per = ctx.d, rv.per
    n = per.indice
    cur_m = d.leads_mensuels(per.annee)
    old_m = d.leads_mensuels(per.annee - 1)
    cur_vn = d.leads_mensuels(per.annee, projet="VN")
    old_vn = d.leads_mensuels(per.annee - 1, projet="VN")
    if not (sum(cur_m[:n]) or sum(old_m[:n])):
        return None
    pm = rv.prec.indice - 1
    prec_total = d.leads_total(rv.prec)
    prec_vn = d.leads(rv.prec)["VN"]
    f = {
        "n": n, "cur_m": cur_m, "old_m": old_m, "cur_vn": cur_vn, "old_vn": old_vn,
        "mois": cur_m[n - 1], "mois_prec": prec_total, "mois_n1": old_m[n - 1],
        "vn_mois": cur_vn[n - 1], "vn_prec": prec_vn, "vn_n1": old_vn[n - 1],
        "ytd": sum(cur_m[:n]), "ytd_n1": sum(old_m[:n]), "ytd_vn": sum(cur_vn[:n]), "ytd_vn_n1": sum(old_vn[:n]),
    }
    f["e_mois_prec"] = evol(f["mois"], f["mois_prec"])
    f["e_mois_n1"] = evol(f["mois"], f["mois_n1"])
    f["e_vn_n1"] = evol(f["vn_mois"], f["vn_n1"])
    f["e_vn_prec"] = evol(f["vn_mois"], f["vn_prec"])
    f["e_ytd"] = evol(f["ytd"], f["ytd_n1"])
    f["e_ytd_vn"] = evol(f["ytd_vn"], f["ytd_vn_n1"])
    f["part_vn"] = f["ytd_vn"] / f["ytd"] * 100 if f["ytd"] else None
    return f


def _ytd_marques(ctx, rv):
    """[(marque, leads cumulés année, leads cumulés N-1)] décroissant, marques à leads > 0."""
    n = rv.per.indice
    out = []
    for m in ctx.d.marques:
        a = sum(ctx.d.leads_mensuels(rv.per.annee, m)[:n])
        b = sum(ctx.d.leads_mensuels(rv.per.annee - 1, m)[:n])
        if a > 0:
            out.append((m, a, b))
    return sorted(out, key=lambda x: -x[1])


def _nbsp(v):
    return nb(v).replace(" ", "\u00a0")


def _liste(noms):
    noms = list(noms)
    return noms[0] if len(noms) == 1 else ", ".join(noms[:-1]) + " and " + noms[-1]


# ------------------------------------------------------------------ résumé exécutif (diapositive 4)
def _ventes_nsc(ctx, rv):
    """(commandes NSC cumulées de l'année, celles du même cumul N-1, [marques contributrices par écart décroissant]) ; None si la donnée C1ST manque."""
    class _Cumul:
        def __init__(self, annee, mois_fin):
            self._m = [(annee, m) for m in range(1, mois_fin + 1)]

        def mois(self):
            return self._m
    n = rv.per.indice
    a_, b_ = _Cumul(rv.per.annee, n), _Cumul(rv.per.annee - 1, n)
    lignes = []
    for m in ctx.d.marques:
        if m == "SPOTICAR":
            continue
        a, b = ctx.d.c1st_nsc(a_, m), ctx.d.c1st_nsc(b_, m)
        if a is None or b is None:
            continue
        lignes.append((m, a, b))
    if not lignes:
        return None
    ta, tb = sum(x[1] for x in lignes), sum(x[2] for x in lignes)
    contrib = [ctx.nom(m) for m, a, b in sorted(lignes, key=lambda x: -(x[1] - x[2])) if a - b > 0][:3]
    return ta, tb, contrib


def _source_nc(ctx, rv):
    """(part du site de marque dans les leads New cars du mois, source principale, part, [emplacements du site principaux]) ou None."""
    d, per = ctx.d, rv.per
    try:
        if not d.acquisition_disponible(per):
            return None
    except Exception:
        return None
    par_source, total = {}, 0
    supports = {}
    for m in d.marques:
        for src, v in d.sources(per, m).items():
            par_source[src] = par_source.get(src, 0) + v.get("VN", 0)
            total += v.get("VN", 0)
        for sup, v in d.cta(per, m).items():
            if v.get("VN", 0) and not sup.startswith("("):
                supports[sup] = supports.get(sup, 0) + v["VN"]
    if not total:
        return None
    top = max(par_source, key=par_source.get)
    haut = [s for s, _ in sorted(supports.items(), key=lambda x: -x[1])][:3]
    return par_source.get("Main-Website", 0) / total * 100, top, par_source[top] / total * 100, haut


def _puces_resume(ctx, rv, f):
    an, mois = rv.per.annee, _mois_nom(rv.per)
    gras = lambda t: (t, True)
    txt = lambda t: (t, False)
    out = []
    adj = "Strong" if (f["e_ytd"] or 0) >= 10 else "Solid" if (f["e_ytd"] or 0) >= 0 else "Weak"
    nc_dyn = (f["e_ytd_vn"] is not None and f["e_ytd"] is not None and f["e_ytd_vn"] > f["e_ytd"])
    out.append([gras(f"{adj} YTD performance:"), txt(f" {_k(f['ytd'])} leads generated, "), gras(f"{pct(f['e_ytd'], 0, signe=True)} vs {an - 1}") if f["e_ytd"] is not None else txt("n/a vs " + str(an - 1)),
                txt(f", with NC leads at {pct(f['part_vn'])} of the total" + (f" and growing faster ({pct(f['e_ytd_vn'], 0, signe=True)}) than the other leads." if nc_dyn else
                    (f" ({pct(f['e_ytd_vn'], 0, signe=True)} vs {an - 1})." if f["e_ytd_vn"] is not None else ".")))])
    adj = "solid" if (f["e_mois_n1"] or 0) >= 0 else "softer"
    out.append([gras(f"{mois} {adj} vs last year:"), txt(f" {_k(f['mois'])} total leads, "), gras(f"{pct(f['e_mois_n1'], 0, signe=True)} YoY"),
                txt(f" and {pct(f['e_mois_prec'], 0, signe=True)} vs {_mois_nom(rv.prec)}, while NC leads reached {_k(f['vn_mois'])} ("),
                gras(f"{pct(f['e_vn_n1'], 0, signe=True)} YoY"), txt(").")])
    # contraste entre marques : évolution du mois vs N-1, bases d'au moins 100 leads
    ev = []
    for m in MARQUES:
        if m not in ctx.d.marques:
            continue
        a, b = ctx.d.leads_total(rv.per, m), ctx.d.leads_total(rv.n1, m)
        e = evol(a, b)
        if e is not None and b >= 100:
            ev.append((m, e))
    haut = [ctx.nom(m) for m, e in sorted(ev, key=lambda x: -x[1]) if e >= 20][:4]
    bas = [ctx.nom(m) for m, e in sorted(ev, key=lambda x: x[1]) if e < 0][:3]
    if haut or bas:
        seg = [gras("Performance remains contrasted across brands:")]
        if haut and bas:
            seg += [txt(f" {_liste(haut)} recorded strong YoY growth, while "), gras(f"{_liste(bas)} {'is' if len(bas) == 1 else 'are'} declining YoY"), txt(".")]
        elif haut:
            seg += [txt(f" {_liste(haut)} recorded strong YoY growth and no brand is declining.")]
        else:
            seg += [txt(" no brand shows strong YoY growth, while "), gras(f"{_liste(bas)} {'is' if len(bas) == 1 else 'are'} declining YoY"), txt(".")]
        out.append(seg)
    s = _source_nc(ctx, rv)
    if s:
        part_site, top, part_top, haut_sup = s
        if top == "Main-Website":
            seg = [gras("Main brand websites are the key source of NC leads"), txt(f" ({pct(part_site)} of NC leads in {mois}")]
            seg.append(txt(f"), with {_liste(haut_sup)} emerging as the most effective onsite entry points." if haut_sup else ")."))
        else:
            seg = [gras(f"{top} is the key source of NC leads"), txt(f" ({pct(part_top)} of NC leads in {mois}), ahead of main brand websites ({pct(part_site)})."
                                                                   + (f" Onsite, {_liste(haut_sup)} lead the entry points." if haut_sup else ""))]
        out.append(seg)
    v = _ventes_nsc(ctx, rv)
    if v:
        ta, tb, contrib = v
        diff = ta - tb
        seg = [gras("Lead growth is translating into additional sales:" if diff > 0 else "Additional sales are under pressure:"),
               txt(f" {_nbsp(ta)} NC sales have already been generated from {an} leads in the period reported, "),
               gras(f"{_nbsp(abs(diff))} {'more' if diff >= 0 else 'fewer'} than last year")]
        seg.append(txt(f", with {_liste(contrib)} identified as the main contributor{'s' if len(contrib) > 1 else ''}" if contrib and diff > 0 else ""))
        out.append(seg)
    return out


def resume_executif(ctx, rv):
    """Diapositive 4 : cinq puces en texte plein écran (16 pt, début en gras)."""
    f = _faits(ctx, rv)
    if not f:
        return None
    s = slide_revue(ctx, rv, "Executive Summary", source=None)
    tb = s.shapes.add_textbox(Inches(0.0), Inches(1.13), Inches(12.94), Inches(3.69))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, runs in enumerate(_puces_resume(ctx, rv, f)):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.5
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", "342900"); pPr.set("indent", "-342900")
        bu_font = etree.SubElement(pPr, qn("a:buFont")); bu_font.set("typeface", "Arial")
        bu = etree.SubElement(pPr, qn("a:buChar")); bu.set("char", "•")
        for texte, gras in runs:
            if not texte:
                continue
            r = p.add_run(); r.text = texte
            r.font.size = Pt(16); r.font.bold = gras
    return s


# ------------------------------------------------------------------ vue globale (diapositive 5)
def _bandeau(s, x, y, w, h, texte, fond, couleur="000000", taille=9, gras=True):
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.shadow.inherit = False
    b.fill.solid(); b.fill.fore_color.rgb = g.rgb(fond); b.line.fill.background()
    tf = b.text_frame
    tf.margin_left = Inches(0.05); tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = texte
    r.font.size = Pt(taille); r.font.bold = gras; r.font.color.rgb = g.rgb(couleur)
    return b


def barres_marques(slide, x, y, w, h, lignes, taille=8):
    """Barres horizontales natives (une couleur par marque), la plus grande en haut ; étiquette = part en % (gras) puis valeur (italique gris).
    lignes = [(nom, valeur, part %, couleur)]."""
    cd = CategoryChartData()
    cd.categories = [l[0] for l in lignes]
    cd.add_series("Leads", [l[1] for l in lignes])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_legend = False; ch.has_title = False
    ch.font.size = Pt(taille)
    ch.category_axis.reverse_order = True
    ch.category_axis.format.line.color.rgb = g.rgb("BFBFBF")
    ch.category_axis.tick_labels.font.size = Pt(taille)
    ch.category_axis.tick_labels.font.color.rgb = g.rgb("7F6F5F")
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.value_axis.maximum_scale = max(l[1] for l in lignes) * 1.28
    ch.value_axis.minimum_scale = 0
    plot = ch.plots[0]
    plot.gap_width = 25
    ser = plot.series[0]
    for i, (nom, val, part, coul) in enumerate(lignes):
        pt = ser.points[i]
        pt.format.fill.solid(); pt.format.fill.fore_color.rgb = g.rgb(coul)
        dl = pt.data_label
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        tf = dl.text_frame
        p = tf.paragraphs[0]
        r = p.add_run(); r.text = f"{part:.0f}%  "; r.font.size = Pt(taille); r.font.bold = True; r.font.color.rgb = g.rgb("000000")
        r = p.add_run(); r.text = nb(val); r.font.size = Pt(taille - 0.5); r.font.italic = True; r.font.color.rgb = g.rgb("7F7F7F")
    return gf


def _bloc_chiffre(s, x, titre, libelle, valeur, lib_evol, evolution, w, dy=0.0):
    _bandeau(s, x, 1.43 + dy, w, 0.27, titre, GRIS_BANDEAU, taille=8)
    g.texte(s, x, 1.78 + dy, w, 0.25, libelle, taille=8, couleur="000000")
    g.texte(s, x, 2.22 + dy, w, 0.3, nb(valeur), taille=9, couleur="000000", align=PP_ALIGN.RIGHT)
    g.texte(s, x, 2.92 + dy, w, 0.25, lib_evol, taille=7, couleur="000000")
    coul = "7F7F7F" if evolution is None else ("C00000" if evolution < 0 else "2E7D32")
    g.texte(s, x, 3.3 + dy, w, 0.3, pct(evolution, 0, signe=True) if evolution is not None else "n/a", taille=9, gras=True, couleur=coul, align=PP_ALIGN.RIGHT)


def _mois_glissants(per, k=14):
    """Les k derniers mois (année, mois) se terminant au mois de la période."""
    n = per.annee * 12 + per.indice - 1
    return [divmod(n - i, 12) for i in range(k - 1, -1, -1)]


def _legende_projets(s, x, y):
    cadre = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(1.45), Inches(0.78))
    cadre.shadow.inherit = False
    cadre.fill.solid(); cadre.fill.fore_color.rgb = g.rgb("FFFFFF"); cadre.line.color.rgb = g.rgb("D9D9D9")
    g.texte(s, x + 0.03, y + 0.02, 1.4, 0.2, "Purchase Project", taille=8, couleur="404040")
    for i, (nom, coul) in enumerate((("New cars & Demo cars", C_VN), ("Trade-in only", C_SANS), ("Used cars", C_VO))):
        yy = y + 0.25 + i * 0.165
        c = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.07), Inches(yy + 0.03), Inches(0.1), Inches(0.1))
        c.shadow.inherit = False
        c.fill.solid(); c.fill.fore_color.rgb = g.rgb(coul); c.line.fill.background()
        g.texte(s, x + 0.2, yy - 0.02, 1.25, 0.2, nom, taille=8, couleur="000000")


def _puces_global(ctx, rv, f, ytd_m):
    """Les trois commentaires de la vue « Global Overview » (réutilisés tels quels par moderne.py)."""
    per, an = rv.per, rv.per.annee
    top2 = [ctx.nom(m) for m, _, _ in ytd_m[:2]]
    part_top2 = sum(a for _, a, _ in ytd_m[:2]) / f["ytd"] * 100 if f["ytd"] else None
    puces = [f"{_mois_nom(per)} reached {_k(f['mois'])} total leads, {pct(f['e_mois_prec'], 0, signe=True)} MoM and {pct(f['e_mois_n1'], 0, signe=True)} YoY.",
             f"NC leads: {_k(f['vn_mois'])} leads in {_mois_nom(per)}, {pct(f['e_vn_n1'], 0, signe=True)} YoY and {pct(f['e_vn_prec'], 0, signe=True)} vs {_mois_nom(rv.prec)}.",
             f"YTD performance at {_k(f['ytd'])} leads, {pct(f['e_ytd'], 0, signe=True)} vs {an - 1}" +
             (f", with {_liste(top2)} accounting for {pct(part_top2)} of generated leads." if len(top2) == 2 else ".")]
    return puces


def vue_globale(ctx, rv):
    """Diapositive 5 : tableau de bord « EXEC SUMMARY » (chiffres du mois et du cumul, barres par marque, tendance par projet d'achat) + commentaires."""
    f = _faits(ctx, rv)
    if not f:
        return None
    per, an = rv.per, rv.per.annee
    mois_maj = MOIS_COURTS["en"][per.indice - 1].upper()
    ytd_m = _ytd_marques(ctx, rv)
    puces = _puces_global(ctx, rv, f, ytd_m)
    s = slide_revue(ctx, rv, "Global Overview", puces=puces, boite_pleine=True)
    _bandeau(s, 0.69, 1.14, 9.87, 0.26, f"EXEC SUMMARY - TRADE IN SITUATION {mois_label(per).upper()}", NAVY_TABLEAU, couleur="FFFFFF", taille=9)
    _bloc_chiffre(s, 0.69, "REPORTING MONTHLY", f"Total generated leads {mois_maj}. {an}", f["mois"], "Evolution vs last month", f["e_mois_prec"], 1.83)
    _bloc_chiffre(s, 2.58, "REPORTING YTD", f"Total generated leads {an}", f["ytd"], f"Evolution vs last year ( {an - 1} et {an})", f["e_ytd"], 1.75)
    _bandeau(s, 4.38, 1.43, 6.18, 0.27, f"Generated leads YTD {mois_label(per).upper()} : split by brand", GRIS_BANDEAU, taille=8)
    if ytd_m:
        tot = f["ytd"] or 1
        lignes = [("S&Y" if m == "STELLANTIS &YOU" else m.replace(" ", "_"), a, a / tot * 100, COULEUR_MARQUE.get(m, "BAB0AC")) for m, a, _ in ytd_m]
        barres_marques(s, 4.38, 1.72, 6.18, 2.12, lignes)
    _bandeau(s, 0.69, 3.87, 9.87, 0.26, "Trend YTD", GRIS_BANDEAU, taille=8)
    mois = _mois_glissants(per)
    cache = {}
    def serie(projet):
        out = []
        for a, m0 in mois:
            if a not in cache:
                cache[a] = {p: ctx.d.leads_mensuels(a, projet=p) for p in ("VN", "SANS", "VO")}
            out.append(cache[a][projet][m0])
        return out
    cats = [f"{MOIS_COURTS['en'][m0]} {a}" for a, m0 in mois]
    gf = g.lignes_chart(s, 0.69, 4.15, 9.87, 2.7, cats, [("New cars & Demo cars", serie("VN"), C_VN), ("Trade-in only", serie("SANS"), C_SANS), ("Used cars", serie("VO"), C_VO)],
                   taille=8, legende=False, etiquettes=True)
    try:                                   # étiquettes : Used cars sous la courbe (évite les chevauchements avec New cars)
        gf.chart.plots[0].series[2].data_labels.position = XL_LABEL_POSITION.BELOW
    except Exception:
        pass
    _legende_projets(s, 11.41, 5.92)
    return s


# ------------------------------------------------------------------ tendance des leads (diapositive 6)
def _puces_tendance(ctx, rv, f):
    """Les trois commentaires de la diapositive « Lead trend » (réutilisés tels quels par la présentation au nouveau design, moderne.py)."""
    per, an, n = rv.per, rv.per.annee, rv.per.indice
    audessus = sum(1 for i in range(n) if f["cur_m"][i] > f["old_m"][i])
    mois, prec = _mois_nom(per), _mois_nom(rv.prec)
    e_prec_n1 = evol(f["cur_m"][rv.prec.indice - 1], f["old_m"][rv.prec.indice - 1]) if n > 1 else None
    e_vn_prec_n1 = evol(f["cur_vn"][rv.prec.indice - 1], f["old_vn"][rv.prec.indice - 1]) if n > 1 else None
    if audessus == n:
        p1 = f"Total leads remained above {an - 1} throughout the year, with growth at {pct(f['e_mois_n1'], 0, signe=True)} YoY in {mois}" + (f" vs {pct(e_prec_n1, 0, signe=True)} in {prec}." if e_prec_n1 is not None else ".")
    else:
        p1 = f"Total leads were above {an - 1} in {audessus} of the {n} months so far, at {pct(f['e_mois_n1'], 0, signe=True)} YoY in {mois}" + (f" vs {pct(e_prec_n1, 0, signe=True)} in {prec}." if e_prec_n1 is not None else ".")
    sens = "slowdown" if (e_vn_prec_n1 is not None and f["e_vn_n1"] is not None and f["e_vn_n1"] < e_vn_prec_n1) else "acceleration"
    p2 = f"NC leads {'remain the main growth driver, ' if (f['e_vn_n1'] or 0) > (f['e_mois_n1'] or 0) else ''}{pct(f['e_vn_n1'], 0, signe=True)} YoY in {mois}" + \
         (f", with {'an' if sens[0] == 'a' else 'a'} {sens} vs {prec} ({pct(e_vn_prec_n1, 0, signe=True)})." if e_vn_prec_n1 is not None else ".")
    p3 = f"{mois} volumes: {nb(f['mois'])} total leads ({pct(f['e_mois_prec'], 0, signe=True)} vs {prec}) and {nb(f['vn_mois'])} NC leads ({pct(f['e_vn_prec'], 0, signe=True)} vs {prec})."
    return [p1, p2, p3]


def tendance_leads(ctx, rv):
    """Diapositive 6 : barres mensuelles N-1 / N, total des leads puis leads New cars, toutes marques du périmètre."""
    f = _faits(ctx, rv)
    if not f:
        return None
    per, an, n = rv.per, rv.per.annee, rv.per.indice
    ld = ctx.d.dernier_jour_leads()
    ecoules = n if (ld and ld >= per.debut.isoformat()) else n - 1
    p1, p2, p3 = _puces_tendance(ctx, rv, f)
    s = slide_revue(ctx, rv, "Executive Summary", sous_titre="Lead trend – all brands", puces=[p1, p2, p3], boite_pleine=True)
    graphique_annuel(s, 0.72, 1.22, 9.67, 2.75, "Total leads year trend", f["old_m"], f["cur_m"], ecoules, an)
    graphique_annuel(s, 0.72, 4.05, 9.67, 2.85, "NC leads year trend", f["old_vn"], f["cur_vn"], ecoules, an)
    return s
