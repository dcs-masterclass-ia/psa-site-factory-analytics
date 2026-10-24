"""Leads back-office de l'espace OEM (Ford) -> data/oem/<slug>-leads.json, fusionné ensuite dans data/oem/<slug>.json par pipeline.oem.

Pourquoi un fichier à part : le back-office Ford (prod-external.autobiz.com) n'est PAS couvert par l'API d'extraction à jeton des sites Stellantis (pipeline/leads_extract.py),
et sa connexion est protégée par un reCAPTCHA : pas de rafraîchissement nocturne possible tant que l'équipe BO n'a pas ouvert une route. La collecte GA4 de la nuit (oem.py)
réécrit `<slug>.json` sans toucher à ce fichier, qu'elle relit pour le fusionner : les leads ne sont donc jamais perdus par la collecte nocturne.

    python -m pipeline.oem_leads --fichier extract.xlsx [--ecrire]        # depuis une extraction téléchargée à la main
    python -m pipeline.oem_leads --cookie-fichier ~/.claude/secrets/bo-ford-cookie [--ecrire]   # génère l'extraction avec une session BO valide

Mêmes règles que côté Stellantis : seuls les leads de production comptent (hors tests, doublons, aperçus). Aucune donnée personnelle n'est conservée : seulement des comptages."""

import argparse
import io
import json
import sys
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

RACINE = Path(__file__).resolve().parent.parent
DATA = RACINE / "data" / "oem"

BO = "https://prod-external.autobiz.com"
SITES = {"ford-fr": {"nom": "FORD FR", "site_id": 14}}

# colonnes de l'extraction Ford (noms différents de ceux des extractions Stellantis)
C_DATE, C_MARQUE, C_ENERGIE, C_PROJET = "CREATION_DATE", "TRADE_IN_BRAND", "FUEL", "PURCHASE_PROJECT_TYPE"
C_SAISIE, C_SOURCE, C_ACQ, C_CAMPAGNE, C_APPAREIL, C_PDV = "ENTRY_TYPE", "UTM_SOURCE", "SOURCE_ACQUISITION", "UTM_CAMPAIGN", "DEVICE_TYPE", "DEALER_NAME"
C_TEST, C_DOUBLON, C_APERCU = "TEST", "DUPLICATE", "PREVIEW"
SAISIE = {"registration": "Avec immatriculation", "brand": "Marque / modele"}
PROJET_VIDE = "No purchase project"


def telecharge(site_id, cookie, debut, fin):
    """Génère l'extraction du back-office (lecture seule : elle ne modifie rien) et la renvoie en octets xlsx."""
    corps = urllib.parse.urlencode({"dateBegin": debut, "dateEnd": fin, "recordId": "", "selectCrm": "-1", "dealerCode": "", "siteId": str(site_id)}).encode()
    req = urllib.request.Request(f"{BO}/admin/generate/extract/{site_id}", data=corps, headers={"Cookie": f"PHPSESSID={cookie}", "User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=300) as r:
        if "spreadsheet" not in (r.headers.get("content-type") or ""):
            raise RuntimeError("le back-office n'a pas renvoyé de fichier Excel (session expirée ?)")
        return r.read()


def lignes_xlsx(source):
    import openpyxl
    wb = openpyxl.load_workbook(source if isinstance(source, (str, Path)) else io.BytesIO(source), read_only=True)
    it = wb[wb.sheetnames[0]].iter_rows(values_only=True)
    entetes = [str(c) if c is not None else "" for c in next(it)]
    for r in it:
        yield dict(zip(entetes, r))


def _jour(v):
    if isinstance(v, datetime):
        return v.date()
    try:
        return date.fromisoformat(str(v)[:10])
    except ValueError:
        return None


def valide(l):
    return l.get(C_TEST) == "NO" and l.get(C_DOUBLON) == "NO" and l.get(C_APERCU) == "NO"


def _top(c, n=20):
    return [[k, v] for k, v in c.most_common(n) if k]


def construit(lignes, aujourdhui=None):
    """leads[mois] au format des fichiers de site Stellantis (total, daily, brand, fuel, entry, project, source, code, device, dailyDevice, pdv) + `total` (daily = mois mis bout à bout)."""
    aujourdhui = aujourdhui or date.today()
    par_mois = defaultdict(list)
    for l in lignes:
        j = _jour(l.get(C_DATE))
        if j and valide(l):
            par_mois[f"{j.year}-{j.month:02d}"].append((j, l))
    if not par_mois:
        return {}, []
    mois = sorted(par_mois)
    # tous les mois entre le premier et le dernier, même sans lead
    a, m = int(mois[0][:4]), int(mois[0][5:])
    z = (int(mois[-1][:4]), int(mois[-1][5:]))
    tous = []
    while (a, m) <= z:
        tous.append(f"{a}-{m:02d}")
        m += 1
        if m == 13:
            a, m = a + 1, 1
    leads, bout = {}, []
    for mm in tous:
        an, mo = int(mm[:4]), int(mm[5:])
        jours = (date(an + (mo == 12), mo % 12 + 1, 1) - date(an, mo, 1)).days
        bornes = min(jours, aujourdhui.day) if (an, mo) == (aujourdhui.year, aujourdhui.month) else jours
        daily = [0] * bornes
        marque, energie, projet, source, code, pdv, appareil = (Counter() for _ in range(7))
        saisie = {"Avec immatriculation": 0, "Marque / modele": 0}
        dd = defaultdict(lambda: [0] * bornes)
        for j, l in par_mois.get(mm, []):
            if j.day <= bornes:
                daily[j.day - 1] += 1
            for cpt, col in ((marque, C_MARQUE), (energie, C_ENERGIE)):
                if l.get(col) not in (None, "None", ""):
                    cpt[str(l[col])] += 1
            p = l.get(C_PROJET)
            projet[p if p not in (None, "None", "") else PROJET_VIDE] += 1
            s = l.get(C_ACQ) if l.get(C_ACQ) not in (None, "None", "") else l.get(C_SOURCE)
            if s not in (None, "None", ""):
                source[str(s)] += 1
            if l.get(C_CAMPAGNE) not in (None, "None", ""):
                code[str(l[C_CAMPAGNE])] += 1
            if l.get(C_PDV) not in (None, "None", ""):
                pdv[str(l[C_PDV]).strip()] += 1
            ap = l.get(C_APPAREIL)
            if ap not in (None, "None", ""):
                appareil[str(ap)] += 1
                if j.day <= bornes:
                    dd[str(ap)][j.day - 1] += 1
            if l.get(C_SAISIE) in SAISIE:
                saisie[SAISIE[l[C_SAISIE]]] += 1
        leads[mm] = {"total": sum(daily), "daily": daily, "brand": _top(marque), "fuel": _top(energie), "entry": saisie, "project": _top(projet),
                     "source": _top(source), "code": _top(code), "device": _top(appareil), "dailyDevice": dict(dd), "pdv": _top(pdv, 30)}
        bout += daily
    leads["total"] = {"total": sum(bout), "daily": bout}
    return leads, tous


def fichier(slug, leads, mois):
    return {"site": SITES[slug]["nom"], "espace": "oem", "generee": datetime.now().isoformat(timespec="seconds"), "mois": mois, "leads": leads,
            "source": "back-office Ford (prod-external.autobiz.com), leads de production (hors tests, doublons, aperçus) — rafraîchissement manuel"}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--site", default="ford-fr")
    ap.add_argument("--fichier")
    ap.add_argument("--cookie-fichier")
    ap.add_argument("--debut", default="2025-01-01")
    ap.add_argument("--ecrire", action="store_true")
    a = ap.parse_args()
    if a.fichier:
        src = a.fichier
    elif a.cookie_fichier:
        src = telecharge(SITES[a.site]["site_id"], Path(a.cookie_fichier).expanduser().read_text().strip(), a.debut, date.today().isoformat())
    else:
        sys.exit("--fichier ou --cookie-fichier requis")
    leads, mois = construit(lignes_xlsx(src))
    if not leads:
        sys.exit("aucun lead valide dans l'extraction : rien écrit")
    print(f"{SITES[a.site]['nom']} : {leads['total']['total']} leads valides sur {len(mois)} mois ({mois[0]} → {mois[-1]})")
    for mm in mois[-4:]:
        print(f"  {mm} : {leads[mm]['total']}  projets {leads[mm]['project']}")
    if a.ecrire:
        DATA.mkdir(parents=True, exist_ok=True)
        (DATA / f"{a.site}-leads.json").write_text(json.dumps(fichier(a.site, leads, mois), ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
        print(f"écrit : {DATA / (a.site + '-leads.json')}")


if __name__ == "__main__":
    main()
